#!/usr/bin/env python3
# ═══════════════════════════════════════════════════════════════════════════
#  CANDecode — DBC-guided CAN bus log decoder, plotter, and Excel exporter
#
#  An open-source personal project. No proprietary or vendor DBC content is
#  bundled or referenced — bring your own .dbc file(s) and trace log.
#
#  Author  : Dhanoosh B
#  Contact : dhanoosh2001@gmail.com
#  GitHub  : https://github.com/DevDhanoosh
#  License : MIT (see LICENSE)
#  Repo    : this folder is the whole project — a single dependency-light script
# ═══════════════════════════════════════════════════════════════════════════
"""
CANDecode — decode a CAN bus trace against one or more Vector DBC databases,
plot every matched signal, and export a tidy Excel workbook.

Decoding follows the standard Vector DBC convention: Intel (little-endian)
and Motorola (big-endian, "sawtooth") bit layouts, with 2's-complement signed
signals.

Supported trace formats (auto-detected):
  * BUSMASTER ASCII log        (.log / .asc / .txt)
  * Quoted CSV trace           (.csv)  — "Time","ID(hex)","Std/Ext","DLC","Data"
                                          style export, MiniMon-compatible

Typical use
-----------
    python candecode.py --dbc vehicle.dbc --log trace.log --out ./out
    python candecode.py --dbc a.dbc b.dbc --log capture.csv --out ./out \
        --signals MotorTorque PackCurrent --show

Run with no arguments to get interactive file-picker dialogs instead.

Outputs (under --out):
    <log>_decoded.xlsx      Log Info / Summary / CAN Frames / Merged (every
                             decoded signal) / Charts (only the signals you
                             asked to chart, via --signals)
    plots/<signal>.png      one time-series graph per charted signal
    plots/_overlay.png      charted signals overlaid, independent y-axes
    plots/_overlay_norm.png same, each signal normalized 0..1

Dependencies: numpy (required), matplotlib + openpyxl (optional, for plots /
Excel). tqdm is optional — a built-in progress bar is used if it is absent.
"""

import argparse
import os
import re
import sys
import time

try:
    import numpy as np
except ImportError:
    sys.exit("ERROR: numpy is required.  Install it with:  pip install numpy")

# ── optional deps, all degrade gracefully ───────────────────────────────────
try:
    from tqdm import tqdm
except ImportError:                       # minimal drop-in if tqdm isn't present
    class tqdm:                           # noqa: N801
        def __init__(self, iterable=None, total=None, desc="", unit="", **_):
            self.iterable = iterable
            self.total = total if total is not None else (len(iterable) if iterable is not None else 0)
            self.desc = desc
            self.n = 0
            self._last = 0.0
            self._draw()

        def __iter__(self):
            for item in self.iterable:
                yield item
                self.update(1)
            self.close()

        def update(self, k=1):
            self.n += k
            now = time.time()
            if now - self._last > 0.1 or self.n >= self.total:
                self._draw()
                self._last = now

        def _draw(self):
            total = self.total or 1
            frac = min(self.n / total, 1.0)
            bar = "#" * int(frac * 28)
            sys.stderr.write(f"\r  {self.desc:<22} [{bar:<28}] {frac*100:5.1f}%")
            sys.stderr.flush()

        def close(self):
            sys.stderr.write("\n")
            sys.stderr.flush()

        def __enter__(self):
            return self

        def __exit__(self, *_):
            self.close()

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

__version__ = "1.0.0"


# ════════════════════════════════════════════════════════════════════════════
#  DBC parsing
# ════════════════════════════════════════════════════════════════════════════
MESSAGE_LINE_RE = re.compile(r"^BO_\s+(\d+)\s+(\w+)\s*:\s*(\d+)\s+(\S+)")
SIGNAL_LINE_RE = re.compile(
    r'^\s*SG_\s+(\w+)\s*:\s*(\d+)\|(\d+)@(\d)([+-])\s*'
    r'\(([^,]+),([^)]+)\)\s*\[([^|]*)\|([^\]]*)\]\s*"([^"]*)"'
)


def load_message_definitions(text):
    """Return a list of message dicts with nested signal dicts."""
    messages = []
    current = None
    for raw in text.splitlines():
        m = MESSAGE_LINE_RE.match(raw)
        if m:
            mid = int(m.group(1))
            extended = mid >= 0x80000000
            if extended:
                mid -= 0x80000000
            current = dict(id=mid, name=m.group(2), dlc=int(m.group(3)),
                           sender=m.group(4), extended=extended, signals=[])
            messages.append(current)
            continue
        m = SIGNAL_LINE_RE.match(raw)
        if m and current is not None:
            current["signals"].append(dict(
                name=m.group(1), start_bit=int(m.group(2)), length=int(m.group(3)),
                little_endian=(m.group(4) == "1"), signed=(m.group(5) == "-"),
                scale=float(m.group(6)), offset=float(m.group(7)),
                min=float(m.group(8)), max=float(m.group(9)), unit=m.group(10),
            ))
    return messages


BUS_SPEED_LINE_RE = re.compile(r"^BS_\s*:\s*(\d+)")


def read_declared_bus_speed(text):
    """Vector DBC BS_: record, when present, declares the bus speed in
    kbit/s (e.g. 'BS_: 500' -> 500000 bit/s). Most modern DBC exports leave
    this blank — returns None when there's nothing to read."""
    for raw in text.splitlines():
        m = BUS_SPEED_LINE_RE.match(raw.strip())
        if m:
            return float(m.group(1)) * 1000.0
    return None


def declared_bus_speed(dbc_paths):
    """(baudrate_or_None, conflicting) — the bus speed declared by BS_:
    across the given DBCs. If more than one DBC declares a different
    value, conflicting=True and the first non-None value wins."""
    found = []
    for path in dbc_paths:
        with open(path, "r", errors="replace") as fh:
            br = read_declared_bus_speed(fh.read())
        if br is not None:
            found.append(br)
    if not found:
        return None, False
    return found[0], len(set(found)) > 1


def combine_databases(dbc_paths, max_dbc=10):
    """Parse and merge several DBCs; first definition of an ID wins."""
    messages, seen, dups = [], set(), 0
    if len(dbc_paths) > max_dbc:
        print(f"  note: {len(dbc_paths)} DBCs given, using first {max_dbc}.")
        dbc_paths = dbc_paths[:max_dbc]
    for path in dbc_paths:
        with open(path, "r", errors="replace") as fh:
            msgs = load_message_definitions(fh.read())
        added = 0
        for msg in msgs:
            if msg["id"] in seen:
                dups += 1
                continue
            seen.add(msg["id"])
            msg["source"] = os.path.basename(path)
            messages.append(msg)
            added += 1
        print(f"  {os.path.basename(path):<28} {len(msgs):>4} message(s), {added} new")
    if dups:
        print(f"  {dups} duplicate message ID(s) across databases — kept first definition")
    return messages


# ════════════════════════════════════════════════════════════════════════════
#  Log parsing  (BUSMASTER + a generic quoted-CSV trace), auto-detected
# ════════════════════════════════════════════════════════════════════════════
TRACE_LINE_RE = re.compile(
    r"^(\d{2}):(\d{2}):(\d{2}):(\d{4})\s+(?:Rx|Tx)\s+\d+\s+"
    r"0x([0-9A-Fa-f]+)\s+([sx])\s+(\d+)((?:\s+[0-9A-Fa-f]{2})*)"
)


def sniff_trace_format(text_head):
    if '"Identifier (hex)"' in text_head:
        return "csv"
    if re.search(r"\bBUSMASTER\b", text_head, re.I):
        return "busmaster"
    # content sniff: quoted-CSV data rows => csv
    if re.search(r'^\s*"[^"]*"\s*,\s*"[0-9A-Fa-f]+"', text_head, re.M):
        return "csv"
    return "busmaster"


class CanFrame:
    __slots__ = ("t", "id", "dlc", "data", "extended")

    def __init__(self, t, cid, dlc, data, extended):
        self.t = t
        self.id = cid
        self.dlc = dlc
        self.data = data
        self.extended = extended


# ── CAN protocol error-frame recognition ─────────────────────────────────────
# Best-effort: a plain ASCII trace only reports an error frame's sub-type
# (Stuff / Form / CRC / ACK / Bit / Overload) if the logging tool itself wrote
# that word into the line. We recognise it when present and otherwise still
# count the frame under "Error CanFrame (type not logged)" — we never fabricate
# a sub-type the source log didn't actually record.
ERROR_TAGS = [
    ("stuff", "Stuff Error"), ("form", "Form Error"), ("crc", "CRC Error"),
    ("ack", "ACK Error"), ("bit", "Bit Error"), ("overload", "Overload CanFrame"),
]
TRACE_TIMESTAMP_RE = re.compile(r"^(\d{2}):(\d{2}):(\d{2}):(\d{4})\s+(.*)$")


def _label_error_frame(text):
    low = text.lower()
    for kw, label in ERROR_TAGS:
        if kw in low:
            return label
    return "Error CanFrame (type not logged)"


def ingest_trace_log(text):
    frames, std, ext, bad, errors = [], 0, 0, 0, []
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("***"):
            continue
        m = TRACE_LINE_RE.match(line)
        if not m:
            tm = TRACE_TIMESTAMP_RE.match(line)
            if not tm:
                continue
            hh, mm, ss, frac, rest = tm.groups()
            if re.search(r"error", rest, re.I):
                t = int(hh) * 3600 + int(mm) * 60 + int(ss) + int(frac) / 10000.0
                errors.append(dict(t=t, kind=_label_error_frame(rest), raw=line[:120]))
            else:
                bad += 1
            continue
        hh, mm, ss, frac, idhex, typ, dlc, bytestr = m.groups()
        t = int(hh) * 3600 + int(mm) * 60 + int(ss) + int(frac) / 10000.0
        cid = int(idhex, 16)
        dlc = int(dlc)
        data = bytes(int(b, 16) for b in bytestr.split()) if bytestr.strip() else b""
        if len(data) < dlc:
            bad += 1
            continue
        extended = typ == "x"
        std += not extended
        ext += extended
        frames.append(CanFrame(t, cid, dlc, data[:dlc], extended))
    return frames, dict(std=std, ext=ext, malformed=bad), errors


def _split_quoted_row(line):
    out, cur, in_q = [], [], False
    i, n = 0, len(line)
    while i < n:
        c = line[i]
        if in_q:
            if c == '"':
                if i + 1 < n and line[i + 1] == '"':
                    cur.append('"'); i += 1
                else:
                    in_q = False
            else:
                cur.append(c)
        elif c == '"':
            in_q = True
        elif c == ",":
            out.append("".join(cur)); cur = []
        else:
            cur.append(c)
        i += 1
    out.append("".join(cur))
    return out


ROW_TIMESTAMP_RE = re.compile(r"^\s*(\d{1,2}):(\d{2}):(\d{2})[.:](\d{1,6})\s*$")


def ingest_quoted_csv(text):
    """Quoted CSV trace: Time, ID(hex), Std/Ext, DLC, Data(hex bytes)."""
    frames, std, ext, bad, errors = [], 0, 0, 0, []
    started = False
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if not started:
            if re.match(r'^"?\s*Time\s*"?\s*,', line, re.I):
                started = True
            continue
        f = _split_quoted_row(line)
        if len(f) < 5:
            bad += 1
            continue
        tm = ROW_TIMESTAMP_RE.match(f[0])
        if not tm:
            bad += 1
            continue
        frac = tm.group(4)
        t = int(tm.group(1)) * 3600 + int(tm.group(2)) * 60 + int(tm.group(3)) \
            + int(frac) / (10 ** len(frac))
        if re.search(r"\berr", ",".join(f), re.I):
            errors.append(dict(t=t, kind=_label_error_frame(",".join(f)), raw=line[:120]))
            continue
        try:
            cid = int(f[1].strip(), 16)
        except ValueError:
            bad += 1
            continue
        datastr = f[4].strip()
        data = bytes(int(b, 16) for b in datastr.split()) if datastr else b""
        extended = "ext" in f[2].lower()
        std += not extended
        ext += extended
        frames.append(CanFrame(t, cid, len(data), data, extended))
    return frames, dict(std=std, ext=ext, malformed=bad), errors


# ── metadata (date / start / end / duration) ────────────────────────────────
def _normalize_date(raw, hint):
    parts = [p for p in re.split(r"[^0-9]+", raw) if p]
    if len(parts) < 3:
        return raw
    yi = next((i for i, x in enumerate(parts) if len(x) == 4), 2)
    year = parts[yi]
    two = [int(x) for i, x in enumerate(parts) if i != yi][:2]
    if two[0] > 12 >= two[1]:
        d, mo = two
    elif two[1] > 12 >= two[0]:
        mo, d = two
    elif hint == "mdy":
        mo, d = two
    else:
        d, mo = two
    mon = MONTHS[max(0, min(11, mo - 1))]
    return f"{d:02d} {mon} {year}"


def _to_24h(s):
    m = re.search(r"(\d{1,2}):(\d{2})(?::(\d{2}))?\s*(AM|PM)?", s, re.I)
    if not m:
        return s.strip()
    h = int(m.group(1))
    ap = (m.group(4) or "").upper()
    if ap == "PM" and h < 12:
        h += 12
    if ap == "AM" and h == 12:
        h = 0
    return f"{h:02d}:{m.group(2)}:{m.group(3) or '00'}"


def _seconds_to_clock(t):
    t = t % 86400
    h = int(t // 3600); m = int((t % 3600) // 60); s = int(t % 60)
    ms = round((t - int(t)) * 1000)
    return f"{h:02d}:{m:02d}:{s:02d}.{ms:03d}"


def _format_span(sec):
    h = int(sec // 3600); m = int((sec % 3600) // 60); s = sec % 60
    return f"{sec:.3f} s  ({h:02d}:{m:02d}:{s:06.3f})"


def resolve_trim_value(val, log_t0):
    """
    Convert a trim value to seconds-from-log-start. Accepts a plain number
    (already relative seconds) or a clock string 'HH:MM:SS[.mmm]' / 'MM:SS',
    which is interpreted against the log clock (absolute minus log start).
    Returns None for blank input; raises ValueError on garbage.
    """
    if val is None:
        return None
    val = str(val).strip()
    if val == "":
        return None
    if ":" in val:
        parts = [float(p) for p in val.split(":")]
        if len(parts) == 3:
            absolute = parts[0] * 3600 + parts[1] * 60 + parts[2]
        elif len(parts) == 2:
            absolute = parts[0] * 60 + parts[1]
        else:
            absolute = parts[0]
        return absolute - log_t0
    return float(val)


def collect_log_metadata(text, frames, fmt):
    meta = dict(format="Quoted CSV trace" if fmt == "csv" else "BUSMASTER",
                version="—", date="—", start="—", end="—",
                duration="—", dur_sec=0.0, frames=len(frames))
    if frames:
        ts = np.fromiter((f.t for f in frames), dtype=np.float64, count=len(frames))
        first, last = float(ts.min()), float(ts.max())
        dur = last - first
        if dur < 0:
            dur += 86400
        meta["dur_sec"] = dur
        meta["duration"] = _format_span(dur)
        meta["start"] = _seconds_to_clock(first)
        meta["end"] = _seconds_to_clock(last)
    if fmt == "csv":
        d = re.search(r"^\s*Date:\s*([0-9]{1,4}[\/\-.][0-9]{1,2}[\/\-.][0-9]{1,4})",
                      text, re.I | re.M)
        if d:
            meta["date"] = _normalize_date(d.group(1), "mdy")
        s = re.search(r"Start time:\s*([0-9:]+\s*(?:AM|PM)?)", text, re.I)
        e = re.search(r"Stop time:\s*([0-9:]+\s*(?:AM|PM)?)", text, re.I)
        if s:
            meta["start"] = _to_24h(s.group(1))
        if e:
            meta["end"] = _to_24h(e.group(1))
    else:
        d = re.search(r"START DATE(?:\s+AND\s+TIME)?\s*[:\-]?\s*"
                      r"([0-9]{1,2}[:\-\/][0-9]{1,2}[:\-\/][0-9]{2,4})", text, re.I)
        if d:
            meta["date"] = _normalize_date(d.group(1), "dmy")
        v = re.search(r"BUSMASTER\s+Ver\s*([0-9][0-9.]*)", text, re.I)
        if v:
            meta["version"] = v.group(1)
    return meta


def open_trace(path):
    with open(path, "r", errors="replace") as fh:
        text = fh.read()
    fmt = sniff_trace_format(text[:3000])
    frames, counts, errors = (ingest_quoted_csv if fmt == "csv" else ingest_trace_log)(text)
    meta = collect_log_metadata(text, frames, fmt)
    return fmt, frames, counts, meta, errors


# ════════════════════════════════════════════════════════════════════════════
#  Signal decode  (vectorized over frames; validated bit-walk)
# ════════════════════════════════════════════════════════════════════════════
def _bit_layout(sig):
    """List of (byte_idx, bit_in_byte, dest_shift) — exact Vector layout."""
    length = sig["length"]
    out = []
    if sig["little_endian"]:                          # Intel: LSB-first, linear
        for i in range(length):
            N = sig["start_bit"] + i
            out.append((N >> 3, N & 7, i))
    else:                                             # Motorola: MSB-first sawtooth
        bit_pos = sig["start_bit"]
        for i in range(length):
            out.append((bit_pos >> 3, bit_pos & 7, length - 1 - i))
            bit_pos = bit_pos + 15 if (bit_pos & 7) == 0 else bit_pos - 1
    return out


def extract_signal_values(data_matrix, sig):
    """
    data_matrix : uint8 array (n_frames, width) already zero-padded to the
                  message DLC.  Returns float64 array of engineering values.
    """
    length = sig["length"]
    n, width = data_matrix.shape
    if length > 63:                                   # rare wide signal -> Python big-int path
        return _extract_signal_values_bigint(data_matrix, sig)
    raw = np.zeros(n, dtype=np.int64)
    for byte_idx, bit_in_byte, dest in _bit_layout(sig):
        if byte_idx >= width:
            continue
        bit = (data_matrix[:, byte_idx] >> bit_in_byte) & 1
        raw |= bit.astype(np.int64) << dest
    if sig["signed"]:
        sign = np.int64(1) << (length - 1)
        raw = np.where((raw & sign) != 0, raw - (np.int64(1) << length), raw)
    return raw.astype(np.float64) * sig["scale"] + sig["offset"]


def _extract_signal_values_bigint(data_matrix, sig):
    length = sig["length"]
    bmap = _bit_layout(sig)
    out = np.empty(data_matrix.shape[0], dtype=np.float64)
    for r in range(data_matrix.shape[0]):
        row = data_matrix[r]
        raw = 0
        for byte_idx, bit_in_byte, dest in bmap:
            if byte_idx < len(row):
                raw |= int((row[byte_idx] >> bit_in_byte) & 1) << dest
        if sig["signed"] and (raw & (1 << (length - 1))):
            raw -= (1 << length)
        out[r] = raw * sig["scale"] + sig["offset"]
    return out


def build_signal_series(frames, messages, wanted=None):
    """
    Group frames by ID, then decode every (selected) signal on each matched
    message.  Returns list of signal series dicts:
        {name, unit, msg_id, t (ndarray), v (ndarray)}
    """
    frame_ids = set(f.id for f in frames)
    matched_msgs = [m for m in messages if m["id"] in frame_ids]

    # bucket frame indices by id
    buckets = {}
    for i, f in enumerate(frames):
        buckets.setdefault(f.id, []).append(i)

    series = []
    for msg in tqdm(matched_msgs, desc="Decoding", unit="msg"):
        idxs = buckets.get(msg["id"], [])
        if not idxs:
            continue
        width = msg["dlc"]
        # build (n, width) uint8 matrix, zero-padded
        mat = np.zeros((len(idxs), width), dtype=np.uint8)
        t = np.empty(len(idxs), dtype=np.float64)
        for r, fi in enumerate(idxs):
            fr = frames[fi]
            t[r] = fr.t
            b = fr.data[:width]
            mat[r, :len(b)] = np.frombuffer(b, dtype=np.uint8)
        order = np.argsort(t, kind="stable")
        t = t[order]
        mat = mat[order]
        for sig in msg["signals"]:
            if wanted and sig["name"] not in wanted:
                continue
            v = extract_signal_values(mat, sig)
            series.append(dict(name=sig["name"], unit=sig["unit"],
                               msg_id=msg["id"], t=t, v=v))
    return series, matched_msgs


# ════════════════════════════════════════════════════════════════════════════
#  Stats + merge
# ════════════════════════════════════════════════════════════════════════════
def stats(v):
    return dict(n=len(v), min=float(v.min()), max=float(v.max()),
                mean=float(v.mean()), std=float(v.std()))


def nearest_join(times, stimes, svals, tol):
    """Nearest-value join within tolerance; NaN where nothing is close."""
    if len(stimes) == 0:
        return np.full(len(times), np.nan)
    idx = np.clip(np.searchsorted(stimes, times), 0, len(stimes) - 1)
    left = np.clip(idx - 1, 0, len(stimes) - 1)
    d_right = np.abs(stimes[idx] - times)
    d_left = np.abs(stimes[left] - times)
    use_left = d_left < d_right
    best = np.where(use_left, left, idx)
    bestd = np.where(use_left, d_left, d_right)
    return np.where(bestd <= tol, svals[best], np.nan)


# ════════════════════════════════════════════════════════════════════════════
#  Bus load & health — per-ID frequency/jitter/dropout, and an estimated
#  bus-load-over-time curve. Bus load is necessarily an ESTIMATE: a plain
#  ASCII trace doesn't record the exact bit-stuffed frame length, so it's
#  derived from the standard CAN frame overhead (SOF/arbitration/control/
#  CRC/ACK/EOF/IFS) plus a configurable average stuffing inflation factor —
#  good enough to see load trends and hotspots, not a certified measurement.
# ════════════════════════════════════════════════════════════════════════════
def analyze_bus_health(frames, dur_sec, frame_name_by_id, baudrate=500000.0,
                       stuff_factor=1.1, window_sec=1.0, dropout_mult=3.0):
    """
    Returns (per_id, load_ts, load_avg, load_peak):
      per_id  : list of dicts sorted by frame count desc —
                {id, name, count, hz, avg_gap_ms, std_gap_ms, max_gap_ms, dropout}
      load_ts : (bin_starts_sec, load_pct) arrays, one point per `window_sec` bin
      load_avg, load_peak : float, percent
    """
    if not frames or dur_sec <= 0:
        return [], (np.array([]), np.array([])), 0.0, 0.0

    ts = np.fromiter((f.t for f in frames), float, len(frames))
    t0 = float(ts.min())

    by_id = {}
    for f in frames:
        by_id.setdefault(f.id, []).append(f.t)
    per_id = []
    for cid, times in by_id.items():
        times = np.sort(np.asarray(times))
        n = len(times)
        gaps = np.diff(times) * 1000.0                      # ms
        dropout = False
        if len(gaps) >= 4:
            med = float(np.median(gaps))
            if med > 0 and float(gaps.max()) > dropout_mult * med:
                dropout = True
        per_id.append(dict(
            id=cid, name=frame_name_by_id.get(cid, ""), count=n,
            hz=n / dur_sec,
            avg_gap_ms=float(gaps.mean()) if len(gaps) else 0.0,
            std_gap_ms=float(gaps.std()) if len(gaps) else 0.0,
            max_gap_ms=float(gaps.max()) if len(gaps) else 0.0,
            dropout=dropout,
        ))
    per_id.sort(key=lambda r: -r["count"])

    def _frame_bits(f):
        # approx overhead incl. arbitration/control/CRC/ACK/EOF/IFS, before stuffing
        base = 67 if f.extended else 47
        return (base + 8 * f.dlc) * stuff_factor

    n_bins = max(1, int(np.ceil(dur_sec / window_sec)))
    bin_bits = np.zeros(n_bins)
    for f in frames:
        idx = min(n_bins - 1, max(0, int((f.t - t0) / window_sec)))
        bin_bits[idx] += _frame_bits(f)
    # The last bin usually covers less than a full window_sec of real time
    # (dur_sec rarely divides evenly) — dividing it by a full window's
    # capacity like every other bin understates its load%, so give it its
    # own (shorter) duration.
    bin_durs = np.full(n_bins, window_sec)
    bin_durs[-1] = dur_sec - (n_bins - 1) * window_sec
    load_pct = bin_bits / (bin_durs * baudrate) * 100.0
    bin_starts = np.arange(n_bins) * window_sec
    return per_id, (bin_starts, load_pct), float(load_pct.mean()), float(load_pct.max())


# ════════════════════════════════════════════════════════════════════════════
#  Plots  ("progress graphs")
# ════════════════════════════════════════════════════════════════════════════
PALETTE = ["#d98a00", "#0e9bb8", "#b5259e", "#2aa14a",
           "#1f77d8", "#e8433d", "#7b5cff", "#8a8f00"]


def _thin(t, v, max_n=6000):
    if len(t) <= max_n:
        return t, v
    step = int(np.ceil(len(t) / max_n))
    return t[::step], v[::step]


def render_signal_plots(series, outdir, t0, show=False):
    """Render one PNG per signal in `series`, plus overlay PNGs. `series`
    should already be restricted to whatever the caller wants charted —
    Summary/Merged sheets are populated from the full decoded set elsewhere,
    independently of what gets charted here."""
    try:
        import matplotlib
        if not show:
            matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        print("  matplotlib not installed — skipping plots (pip install matplotlib)")
        return []

    plotdir = os.path.join(outdir, "plots")
    os.makedirs(plotdir, exist_ok=True)
    paths = []

    # one graph per signal
    for s in tqdm(series, desc="Plotting", unit="sig"):
        col = PALETTE[len(paths) % len(PALETTE)]
        t, v = _thin(s["t"] - t0, s["v"])
        fig, ax = plt.subplots(figsize=(9, 3.2), dpi=110)
        ax.plot(t, v, color=col, lw=1.1)
        st = stats(s["v"])
        title = s["name"] + (f"  [{s['unit']}]" if s["unit"] else "")
        ax.set_title(title, fontsize=11, fontweight="bold")
        ax.set_xlabel("t (s)", fontsize=9)
        ax.set_ylabel(s["unit"] or "value", fontsize=9)
        ax.grid(True, alpha=0.25)
        ax.text(0.99, 0.02,
                f"min {st['min']:.3f}   max {st['max']:.3f}   avg {st['mean']:.3f}",
                transform=ax.transAxes, ha="right", va="bottom",
                fontsize=8, color="#555")
        fig.tight_layout()
        safe = re.sub(r"[^\w.-]", "_", s["name"])
        p = os.path.join(plotdir, f"{safe}.png")
        fig.savefig(p)
        plt.close(fig)
        paths.append(p)

    # overlay (independent y-axes, up to first 6 signals) + normalized overlay
    overlay = series[:6]
    if len(overlay) >= 2:
        fig, ax0 = plt.subplots(figsize=(11, 4.2), dpi=110)
        axes = [ax0]
        for i, s in enumerate(overlay):
            ax = ax0 if i == 0 else ax0.twinx()
            if i >= 2:
                ax.spines["right"].set_position(("outward", 46 * (i - 1)))
            col = PALETTE[i % len(PALETTE)]
            t, v = _thin(s["t"] - t0, s["v"])
            ax.plot(t, v, color=col, lw=1.1,
                    label=s["name"] + (f" [{s['unit']}]" if s["unit"] else ""))
            ax.set_ylabel(s["unit"] or s["name"], color=col, fontsize=9)
            ax.tick_params(axis="y", labelcolor=col)
            axes.append(ax)
        ax0.set_xlabel("t (s)")
        ax0.set_title("Signal overlay", fontsize=11, fontweight="bold")
        ax0.grid(True, alpha=0.25)
        lines = [ln for ax in axes for ln in ax.get_lines()]
        ax0.legend(lines, [ln.get_label() for ln in lines], fontsize=8, loc="upper right")
        fig.tight_layout()
        p = os.path.join(plotdir, "_overlay.png")
        fig.savefig(p); plt.close(fig); paths.append(p)

        fig, ax = plt.subplots(figsize=(11, 4.2), dpi=110)
        for i, s in enumerate(overlay):
            col = PALETTE[i % len(PALETTE)]
            t, v = _thin(s["t"] - t0, s["v"])
            rng = (v.max() - v.min()) or 1.0
            ax.plot(t, (v - v.min()) / rng, color=col, lw=1.1, label=s["name"])
        ax.set_xlabel("t (s)"); ax.set_ylabel("normalized 0–1")
        ax.set_title("Signal overlay (normalized)", fontsize=11, fontweight="bold")
        ax.grid(True, alpha=0.25); ax.legend(fontsize=8, loc="upper right")
        fig.tight_layout()
        p = os.path.join(plotdir, "_overlay_norm.png")
        fig.savefig(p); plt.close(fig); paths.append(p)

    if show:
        import matplotlib.pyplot as plt
        if len(overlay) >= 2:
            plt.figure(figsize=(11, 4.5))
            for i, s in enumerate(overlay):
                t, v = _thin(s["t"] - t0, s["v"])
                rng = (v.max() - v.min()) or 1.0
                plt.plot(t, (v - v.min()) / rng, lw=1.1, label=s["name"],
                         color=PALETTE[i % len(PALETTE)])
            plt.legend(); plt.xlabel("t (s)"); plt.ylabel("normalized")
            plt.title("Signal overlay (normalized)"); plt.grid(alpha=0.25)
            plt.tight_layout(); plt.show()
    return paths


def render_bus_load_chart(load_ts, outdir, baudrate, show=False):
    """Bus-load-over-time PNG for the Bus Health sheet. Returns the path, or
    None if matplotlib is unavailable or there's nothing to plot."""
    bin_starts, load_pct = load_ts
    if len(bin_starts) == 0:
        return None
    try:
        import matplotlib
        if not show:
            matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return None

    plotdir = os.path.join(outdir, "plots")
    os.makedirs(plotdir, exist_ok=True)
    fig, ax = plt.subplots(figsize=(9, 3.2), dpi=110)
    ax.plot(bin_starts, load_pct, color="#d9483b", lw=1.2)
    ax.fill_between(bin_starts, load_pct, color="#d9483b", alpha=0.12)
    ax.axhline(float(load_pct.mean()), color="#555", lw=0.8, ls="--")
    ax.set_xlabel("t (s)", fontsize=9)
    ax.set_ylabel("Bus load (%, est.)", fontsize=9)
    ax.set_title(f"Bus load over time — estimated at {baudrate/1000:.0f} kbit/s",
                fontsize=11, fontweight="bold")
    ax.grid(True, alpha=0.25)
    fig.tight_layout()
    p = os.path.join(plotdir, "_bus_load.png")
    fig.savefig(p)
    if not show:
        plt.close(fig)
    return p


# ════════════════════════════════════════════════════════════════════════════
#  Excel export — Log Info / Summary / CAN Frames / Merged / Charts
#  (no per-signal sheets — every decoded signal lives as one column in the
#   Merged sheet instead; Charts only embeds the signals you asked to plot)
# ════════════════════════════════════════════════════════════════════════════
def _unique_sheet_name(name, used):
    n = re.sub(r"[:\\/?*\[\]]", " ", name).strip()[:31] or "Signal"
    out, i = n, 2
    while out.lower() in used:
        out = (n[:28] + "_" + str(i))[:31]
        i += 1
    used.add(out.lower())
    return out


def write_report(path, meta, dbc_names, series, frames, frame_name_by_id,
                 tol, plot_paths, t0, sections=None, progress=None, bus_health=None):
    """
    series   : every decoded signal — feeds Log Info / Summary / Merged.
    plot_paths : PNGs already rendered for whichever signals should appear
               on the Charts sheet (a subset of `series`, chosen by the
               caller — see --signals on the CLI).
    sections : optional set of {'summary','frames','merged','charts','bus_health'}
               controlling which sheets are written (Log Info is always
               included).
    bus_health : optional dict from analyze_bus_health() plus 'errors' (list
               from open_trace), 'baudrate', 'stuff_factor', 'chart' (PNG path)
               -> written to a 'Bus Health' sheet.
    progress : optional callable(fraction: float, label: str) for UI progress.
    """
    def _p(frac, label=""):
        if progress:
            try:
                progress(frac, label)
            except Exception:
                pass

    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Font
    except ImportError:
        print("  openpyxl not installed — skipping Excel (pip install openpyxl)")
        return None

    if sections is None:
        sections = {"summary", "frames", "merged", "charts", "bus_health"}
    bold = Font(bold=True)
    _p(0.02, "Writing header…")

    wb = Workbook()
    used = set()

    # Log Info (always)
    ws = wb.active
    ws.title = _unique_sheet_name("Log Info", used)
    info = [
        ("Format", meta["format"]), ("Version", meta["version"]),
        ("Date", meta["date"]), ("Start time", meta["start"]),
        ("End time", meta["end"]), ("Duration", meta["duration"]),
        ("Trim window", meta.get("trim", "full log")),
        ("Total frames", len(frames)),
        ("DBC databases", ", ".join(dbc_names)),
        ("Message IDs matched",
         ", ".join(sorted({f"0x{s['msg_id']:X}" for s in series}))),
        ("Signals decoded", len(series)),
        ("Generated by", f"CANDecode v{__version__} (open source)"),
    ]
    for k, v in info:
        ws.append([k, v])
    for row in ws.iter_rows(min_col=1, max_col=1):
        row[0].font = bold
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 46

    # Bus Health — per-ID frequency/jitter/dropout, estimated bus load, errors
    if bus_health and "bus_health" in sections:
        ws = wb.create_sheet(_unique_sheet_name("Bus Health", used))
        errors = bus_health.get("errors") or []
        err_types = {}
        for e in errors:
            err_types[e["kind"]] = err_types.get(e["kind"], 0) + 1

        warn_font = Font(bold=True, color="FFCC0000")
        r = 1
        for w in (bus_health.get("baudrate_mismatch"), bus_health.get("overload_warning")):
            if w:
                ws.cell(row=r, column=1, value=f"⚠ {w}").font = warn_font
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
                r += 1
        if r > 1:
            r += 1

        summary = [
            ("Baud rate assumed", f"{bus_health['baudrate']/1000:.0f} kbit/s"),
        ]
        if bus_health.get("dbc_baudrate"):
            summary.append(("Baud rate declared by DBC (BS_:)", f"{bus_health['dbc_baudrate']/1000:.0f} kbit/s"))
        summary += [
            ("Bit-stuffing factor assumed", f"{bus_health['stuff_factor']:.2f}x"),
            ("Average bus load (est.)", f"{bus_health['load_avg']:.2f} %"),
            ("Peak bus load (est.)", f"{bus_health['load_peak']:.2f} %"),
            ("Unique message IDs seen", len(bus_health["per_id"])),
            ("Total error frames", len(errors)),
        ]
        for k, v in summary:
            ws.cell(row=r, column=1, value=k).font = bold
            cell = ws.cell(row=r, column=2, value=v)
            if bus_health.get("overload_warning") and k == "Peak bus load (est.)":
                cell.font = warn_font
            r += 1
        r += 1

        ws.cell(row=r, column=1, value="Error CanFrame Types").font = bold
        r += 1
        ws.cell(row=r, column=1, value="Type").font = bold
        ws.cell(row=r, column=2, value="Count").font = bold
        r += 1
        if err_types:
            for kind, n in sorted(err_types.items(), key=lambda kv: -kv[1]):
                ws.cell(row=r, column=1, value=kind)
                ws.cell(row=r, column=2, value=n)
                r += 1
        else:
            ws.cell(row=r, column=1, value="none detected in this trace")
            r += 1
        r += 1

        ws.cell(row=r, column=1, value="Per-Message-ID Timing").font = bold
        r += 1
        hdr = ["CAN_ID", "Message", "Count", "Freq (Hz)", "Avg gap (ms)",
               "Std gap (ms)", "Max gap (ms)", "Dropout?"]
        for c, h in enumerate(hdr, start=1):
            ws.cell(row=r, column=c, value=h).font = bold
        r += 1
        for row_d in bus_health["per_id"]:
            ws.append([f"0x{row_d['id']:X}", row_d["name"], row_d["count"],
                       round(row_d["hz"], 3), round(row_d["avg_gap_ms"], 2),
                       round(row_d["std_gap_ms"], 2), round(row_d["max_gap_ms"], 2),
                       "YES" if row_d["dropout"] else ""])
        for col, w in zip("ABCDEFGH", (12, 24, 10, 11, 13, 13, 13, 10)):
            ws.column_dimensions[col].width = w

        if errors:
            r = ws.max_row + 2
            ws.cell(row=r, column=1, value="Error CanFrame Log").font = bold
            r += 1
            for c, h in enumerate(["Time_s", "Type", "Raw line"], start=1):
                ws.cell(row=r, column=c, value=h).font = bold
            r += 1
            for e in errors[:5000]:
                ws.append([round(e["t"] - t0, 4), e["kind"], e["raw"]])

        chart = bus_health.get("chart")
        if chart and os.path.exists(chart):
            try:
                img = XLImage(chart)
                img.width, img.height = 640, 228
                ws.add_image(img, "J2")
            except Exception:
                pass

    # Summary — every decoded signal
    if "summary" in sections:
        ws = wb.create_sheet(_unique_sheet_name("Summary", used))
        ws.append(["Signal", "Unit", "Samples", "Min", "Max", "Mean", "Std Dev"])
        for cell in ws[1]:
            cell.font = bold
        for s in series:
            st = stats(s["v"])
            ws.append([s["name"], s["unit"], st["n"],
                       round(st["min"], 4), round(st["max"], 4),
                       round(st["mean"], 4), round(st["std"], 4)])

    # CAN Frames (full raw trace)
    if "frames" in sections:
        _p(0.06, "Writing CAN frames…")
        ws = wb.create_sheet(_unique_sheet_name("CAN Frames", used))
        ws.append(["Time_s", "Clock", "CAN_ID", "Type", "Message", "DLC",
                   "Data (hex)", "B0", "B1", "B2", "B3", "B4", "B5", "B6", "B7"])
        for cell in ws[1]:
            cell.font = bold
        ft0 = t0
        nf = max(1, len(frames))
        for idx, f in enumerate(tqdm(frames, desc="Excel: frames", unit="fr")):
            hexbytes = [f"{b:02X}" for b in f.data]
            row = [round(f.t - ft0, 4), _seconds_to_clock(f.t), f"0x{f.id:X}",
                   "Ext" if f.extended else "Std",
                   frame_name_by_id.get(f.id, ""), f.dlc, " ".join(hexbytes)]
            row += hexbytes[:8] + [""] * (8 - len(hexbytes[:8]))
            ws.append(row)
            if progress and idx % 3000 == 0:
                _p(0.06 + 0.44 * idx / nf, "Writing CAN frames…")

    # Merged (as-of, nearest within tol) — every decoded signal, one column each
    if series and "merged" in sections:
        _p(0.52, "Writing merged sheet…")
        all_times = np.unique(np.round(
            np.concatenate([s["t"] - t0 for s in series]), 4))
        ws = wb.create_sheet(_unique_sheet_name("Merged", used))
        ncell = len(all_times) * len(series)
        if ncell > 1_500_000:
            ws.append([f"Merged sheet skipped: {len(all_times)} time points x "
                       f"{len(series)} signals = {ncell:,} cells exceeds the "
                       f"1,500,000-cell limit."])
            ws.append(["Trim the time window, then export again."])
            ws["A1"].font = bold
        else:
            ws.append(["Time_s"] + [f"{s['name']} [{s['unit']}]" if s["unit"]
                                    else s["name"] for s in series])
            for cell in ws[1]:
                cell.font = bold
            cols = []
            for s in tqdm(series, desc="Excel: merge", unit="sig"):
                col = nearest_join(all_times, s["t"] - t0, s["v"], tol)
                cols.append(col)
            cols = np.array(cols)                       # (nsig, ntimes)
            nt = max(1, len(all_times))
            for i in range(len(all_times)):
                row = [round(float(all_times[i]), 4)]
                for c in range(cols.shape[0]):
                    val = cols[c, i]
                    row.append("" if np.isnan(val) else round(float(val), 4))
                ws.append(row)
                if progress and i % 4000 == 0:
                    _p(0.55 + 0.15 * i / nt, "Writing merged sheet…")

    # Charts sheet — embed only the PNGs the caller asked to be plotted
    single = [p for p in plot_paths if not os.path.basename(p).startswith("_")]
    if single and "charts" in sections:
        _p(0.92, "Embedding charts…")
        ws = wb.create_sheet(_unique_sheet_name("Charts", used))
        r = 1
        for p in single:
            try:
                img = XLImage(p)
                img.width, img.height = 640, 228
                ws.add_image(img, f"A{r}")
                r += 13
            except Exception:
                pass

    _p(0.97, "Saving workbook…")
    wb.save(path)
    _p(1.0, "")
    return path


# ════════════════════════════════════════════════════════════════════════════
#  Main
# ════════════════════════════════════════════════════════════════════════════
def interactive_setup(args):
    """
    Fill in any missing --dbc / --log / --out when the script is launched
    without them (e.g. double-clicked or run with no arguments). Tries native
    file-picker dialogs first, then falls back to typed paths.
    """
    print("No files given on the command line — entering interactive mode.\n"
          "  (tip: you can also run it as:  python candecode.py "
          "--dbc a.dbc --log trace.csv --out ./out)\n")

    dbc, log, out = args.dbc, args.log, args.out
    used_dialog = False
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        used_dialog = True

        if not dbc:
            picked = filedialog.askopenfilenames(
                title="Select DBC database file(s) — up to 10",
                filetypes=[("DBC files", "*.dbc"), ("All files", "*.*")])
            dbc = list(picked)
        if not log:
            log = filedialog.askopenfilename(
                title="Select CAN log  (BUSMASTER .log/.asc  or  quoted CSV)",
                filetypes=[("CAN logs", "*.log *.asc *.csv *.txt"), ("All files", "*.*")])
        if not out:
            chosen = filedialog.askdirectory(
                title="Select output folder  (Cancel = ./candecode_out next to the log)")
            out = chosen or None
        root.destroy()
    except Exception:
        used_dialog = False  # headless / no Tk — fall through to text prompts

    if not used_dialog:
        if not dbc:
            raw = input("Path(s) to DBC file(s) (comma-separated): ").strip()
            dbc = [p.strip().strip('"') for p in raw.split(",") if p.strip()]
        if not log:
            log = input("Path to CAN log file: ").strip().strip('"')
        if not out:
            o = input("Output folder [./candecode_out]: ").strip().strip('"')
            out = o or None

    if not out:
        out = os.path.join(os.path.dirname(os.path.abspath(log)) if log else ".", "candecode_out")

    if args.start is None and args.end is None:
        try:
            print("\nOptional trim — press Enter to skip.")
            print("  Enter seconds from start (e.g. 5) or clock time (e.g. 15:27:00).")
            ts = input("  Trim start: ").strip()
            te = input("  Trim end  : ").strip()
            if ts:
                args.start = ts
            if te:
                args.end = te
        except EOFError:
            pass

    args.dbc, args.log, args.out = dbc, log, out
    return args


def main():
    ap = argparse.ArgumentParser(
        description="CANDecode — DBC-guided CAN log decoder (open source).",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    ap.add_argument("--dbc", nargs="+", default=None, help="one or more .dbc files (max 10)")
    ap.add_argument("--log", default=None, help="BUSMASTER .log/.asc or quoted CSV trace")
    ap.add_argument("--out", default=None, help="output directory (default ./candecode_out)")
    ap.add_argument("--signals", nargs="*", default=None,
                    help="only chart/plot these signal names (default: all decoded "
                         "signals are charted). Summary/Merged always cover every "
                         "decoded signal regardless of this flag.")
    ap.add_argument("--tol", type=float, default=0.1,
                    help="merge tolerance in seconds for the Merged sheet")
    ap.add_argument("--start", default=None,
                    help="trim start — seconds from log start (e.g. 5) or clock time (e.g. 15:27:00)")
    ap.add_argument("--end", default=None,
                    help="trim end — seconds from log start (e.g. 30) or clock time (e.g. 15:28:30)")
    ap.add_argument("--no-excel", action="store_true", help="skip the .xlsx export")
    ap.add_argument("--no-plots", action="store_true", help="skip PNG plots")
    ap.add_argument("--no-frames", action="store_true", help="skip the raw CAN Frames sheet")
    ap.add_argument("--no-bus-health", action="store_true",
                    help="skip the Bus Health sheet (per-ID frequency/jitter, estimated "
                         "bus load, error frame count)")
    ap.add_argument("--baudrate", type=float, default=None,
                    help="CAN bus bit rate, bit/s — used to estimate bus load %%. "
                         "Default: read from the DBC's BS_: record if it declares "
                         "one, else 500000")
    ap.add_argument("--stuff-factor", type=float, default=1.1,
                    help="average bit-stuffing inflation applied to the bus-load "
                         "estimate (1.0 = none, 1.2 = worst case)")
    ap.add_argument("--show", action="store_true", help="also show an overlay interactively")
    ap.add_argument("--version", action="version",
                    version=f"CANDecode {__version__}")
    args = ap.parse_args()

    print(f"CANDecode  v{__version__}  (open source)")
    print("=" * 46)

    interactive = not (args.dbc and args.log)
    if interactive:
        args = interactive_setup(args)

    if not args.dbc:
        sys.exit("No DBC file selected — nothing to decode.")
    if not args.log:
        sys.exit("No log file selected — nothing to decode.")
    if args.out is None:
        args.out = "./candecode_out"

    t_start = time.time()
    os.makedirs(args.out, exist_ok=True)

    print("Parsing DBC database(s)…")
    messages = combine_databases(args.dbc)
    total_sigs = sum(len(m["signals"]) for m in messages)
    print(f"  merged: {len(messages)} unique message(s), {total_sigs} signal(s)")

    dbc_baudrate, dbc_baudrate_conflict = declared_bus_speed(args.dbc)
    baudrate_explicit = args.baudrate is not None
    if dbc_baudrate:
        note = "conflicting values across DBCs, using first" if dbc_baudrate_conflict else "BS_:"
        print(f"  DBC declares bus speed: {dbc_baudrate/1000:.0f} kbit/s ({note})")
    if not baudrate_explicit:
        args.baudrate = dbc_baudrate if dbc_baudrate else 500000.0
    baudrate_mismatch = None
    if baudrate_explicit and dbc_baudrate and abs(args.baudrate - dbc_baudrate) > max(1.0, 0.01 * dbc_baudrate):
        baudrate_mismatch = (f"DBC declares {dbc_baudrate/1000:.0f} kbit/s (BS_:) but this analysis "
                             f"is using {args.baudrate/1000:.0f} kbit/s — bus-load numbers below "
                             f"assume the wrong speed unless that's intentional.")
        print(f"  WARNING: {baudrate_mismatch}")
    print()

    print(f"Loading log: {args.log}")
    fmt, frames, counts, meta, errors = open_trace(args.log)
    fmt_label = "Quoted CSV trace" if fmt == "csv" else "BUSMASTER"
    print(f"  detected format : {fmt_label}")
    if not frames:
        sys.exit("  no frames parsed — check the file / format.")
    print(f"  validation      : {len(frames)} frame(s) — "
          f"{counts['std']} std, {counts['ext']} ext, {counts['malformed']} skipped, "
          f"{len(errors)} error frame(s)")
    print(f"  date/duration   : {meta['date']}  |  {meta['duration']}")
    print(f"  window          : {meta['start']} -> {meta['end']}\n")

    # ── optional time trim (seconds from start, or clock time) ──────────────
    log_t0 = min(f.t for f in frames)
    trim_desc = "full log"
    if args.start is not None or args.end is not None:
        try:
            rs = resolve_trim_value(args.start, log_t0)
            re_ = resolve_trim_value(args.end, log_t0)
        except ValueError:
            sys.exit("  invalid trim value — use seconds (e.g. 5) or clock time "
                     "(e.g. 15:27:00).")
        a = 0.0 if rs is None else max(0.0, rs)
        b = meta["dur_sec"] if re_ is None else min(meta["dur_sec"], re_)
        if b <= a:
            sys.exit(f"  invalid trim: end ({b:.3f}s) must be after start ({a:.3f}s).")
        before = len(frames)
        frames = [f for f in frames if a <= (f.t - log_t0) <= b]
        errors = [e for e in errors if a <= (e["t"] - log_t0) <= b]
        if not frames:
            sys.exit(f"  trim window {a:.3f}-{b:.3f}s contains no frames.")
        trim_desc = (f"{a:.3f}-{b:.3f} s  "
                     f"({_seconds_to_clock(log_t0 + a)}–{_seconds_to_clock(log_t0 + b)})")
        print(f"Trim window       : {trim_desc}  "
              f"({len(frames)} of {before} frames kept)\n")
    meta["trim"] = trim_desc

    # decode EVERY signal the DBC(s) and log agree on — the --signals flag
    # only narrows down what gets charted, further down
    series, matched = build_signal_series(frames, messages, None)
    if not series:
        sys.exit("No signals decoded — no DBC message IDs present in the log"
                 + (" / trim window." if trim_desc != "full log" else "."))
    print(f"  matched {len(matched)} message ID(s), decoded {len(series)} signal(s)\n")

    frame_name_by_id = {m["id"]: m["name"] for m in messages}
    t0 = log_t0   # keep seconds-from-log-start as the time origin, even after trimming

    chart_series = series
    if args.signals:
        wanted = set(args.signals)
        chart_series = [s for s in series if s["name"] in wanted]
        if not chart_series:
            print("  note: none of --signals matched a decoded signal — nothing will be charted.\n")

    plot_paths = []
    if not args.no_plots:
        print("Generating graphs…")
        plot_paths = render_signal_plots(chart_series, args.out, t0, show=args.show)
        if plot_paths:
            print(f"  wrote {len([p for p in plot_paths if not os.path.basename(p).startswith('_')])} "
                  f"signal graph(s) + overlays to {os.path.join(args.out, 'plots')}\n")

    bus_health = None
    if not args.no_bus_health:
        print("Analyzing bus load & health…")
        per_id, load_ts, load_avg, load_peak = analyze_bus_health(
            frames, meta["dur_sec"], frame_name_by_id,
            baudrate=args.baudrate, stuff_factor=args.stuff_factor)
        dropouts = sum(1 for r in per_id if r["dropout"])
        print(f"  {len(per_id)} unique ID(s) — est. bus load avg {load_avg:.2f}% / "
              f"peak {load_peak:.2f}%, {dropouts} ID(s) show a dropout-sized gap, "
              f"{len(errors)} error frame(s)")
        overload_warning = None
        if load_peak > 100.0:
            overload_warning = (f"Estimated peak bus load is {load_peak:.0f}%, which is impossible "
                                f"on a real bus — the assumed baud rate ({args.baudrate/1000:.0f} "
                                f"kbit/s) is very likely too low for this trace. Try --baudrate with "
                                f"the next standard speed up (e.g. 250000, 500000, 1000000).")
            print(f"  WARNING: {overload_warning}")
        print()
        bus_chart = None
        if not args.no_plots:
            bus_chart = render_bus_load_chart(load_ts, args.out, args.baudrate, show=args.show)
        bus_health = dict(per_id=per_id, load_avg=load_avg, load_peak=load_peak,
                          baudrate=args.baudrate, stuff_factor=args.stuff_factor,
                          errors=errors, chart=bus_chart,
                          dbc_baudrate=dbc_baudrate, baudrate_mismatch=baudrate_mismatch,
                          overload_warning=overload_warning)

    if not args.no_excel:
        print("Writing Excel workbook…")
        xlsx = os.path.join(args.out, os.path.splitext(os.path.basename(args.log))[0] + "_decoded.xlsx")
        sections = {"summary", "merged", "charts"}
        if not args.no_frames:
            sections.add("frames")
        if bus_health is not None:
            sections.add("bus_health")
        saved = write_report(xlsx, meta, [os.path.basename(p) for p in args.dbc],
                             series, frames, frame_name_by_id, args.tol, plot_paths, t0,
                             sections=sections, bus_health=bus_health)
        if saved:
            size_mb = os.path.getsize(saved) / 1e6
            print(f"  saved {saved}  ({size_mb:.1f} MB)\n")

    print(f"Done in {time.time() - t_start:.1f}s.")


if __name__ == "__main__":
    _flags = ("-h", "--help", "--version")
    _interactive = (("--dbc" not in sys.argv) or ("--log" not in sys.argv)) \
        and not any(f in sys.argv for f in _flags)
    try:
        main()
    except SystemExit as exc:
        if _interactive:
            if exc.code not in (0, None):
                print(exc.code)
        else:
            raise
    except Exception:
        import traceback
        traceback.print_exc()
    finally:
        if _interactive and sys.stdin and sys.stdin.isatty():
            try:
                input("\nPress Enter to close…")
            except EOFError:
                pass
