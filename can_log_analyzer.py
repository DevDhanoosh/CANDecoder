#!/usr/bin/env python3
"""
CAN log analyzer — DBC-guided decode, validation, plots and Excel export.

Supports three trace formats (auto-detected):
  * BUSMASTER  ASCII log   (.log / .asc / .txt)
  * IXXAT MiniMon V3 CSV   (.csv)
  * candump (Linux SocketCAN) — both the `candump -l` log format
    ("(epoch) iface ID#DATA") and the live/`-t`-timestamped console
    format ("iface  ID   [DLC]  b0 b1 ...").

Decoding matches the Vector DBC convention: Intel (little-endian) and
Motorola (big-endian, "sawtooth") bit layouts, with 2's-complement signed
signals. The Motorola walk and sign-extension are the same logic validated
in the browser tool.

Typical use
-----------
    python can_log_analyzer.py --dbc vehicle.dbc --log trace.log --out ./out
    python can_log_analyzer.py --dbc a.dbc b.dbc --log capture.csv --out ./out \
        --signals MotorTorque PackCurrent --show

Outputs (under --out):
    <log>_report.xlsx      Log Info / Summary / CAN Frames / Merged (all signals)
    plots/<signal>.png     one time-series graph per decoded signal
    plots/_overlay.png     all selected signals, independent y-axes
    plots/_overlay_norm.png  same, each signal normalized 0..1

Dependencies: numpy (required), matplotlib + openpyxl (optional, for plots /
Excel). tqdm is optional — a built-in progress bar is used if it is absent.
"""

import argparse
import os
import re
import sys
import time

# A windowed build (or pythonw.exe / a frozen console=False .exe) has no
# console attached, so sys.stdout/stderr/stdin are None — this module and
# its GUI front-end print/write throughout. Redirect to a null sink up
# front instead of guarding every call site.
if sys.stdout is None:
    sys.stdout = open(os.devnull, "w")
if sys.stderr is None:
    sys.stderr = open(os.devnull, "w")

try:
    import numpy as np
except ImportError:
    sys.exit("ERROR: numpy is required.  Install it with:  pip install numpy")

# ── optional deps, all degrade gracefully ───────────────────────────────────
try:
    from tqdm import tqdm
except ImportError:                       # minimal drop-in if tqdm isn't present
    class tqdm:                           # noqa: N801
        def __init__(self, iterable=None, total=None, desc="", unit="", **_):
            self.iterable = iterable
            self.total = total if total is not None else (len(iterable) if iterable is not None else 0)
            self.desc = desc
            self.n = 0
            self._last = 0.0
            self._draw()

        def __iter__(self):
            for item in self.iterable:
                yield item
                self.update(1)
            self.close()

        def update(self, k=1):
            self.n += k
            now = time.time()
            if now - self._last > 0.1 or self.n >= self.total:
                self._draw()
                self._last = now

        def _draw(self):
            total = self.total or 1
            frac = min(self.n / total, 1.0)
            bar = "#" * int(frac * 28)
            sys.stderr.write(f"\r  {self.desc:<22} [{bar:<28}] {frac*100:5.1f}%")
            sys.stderr.flush()

        def close(self):
            sys.stderr.write("\n")
            sys.stderr.flush()

        def __enter__(self):
            return self

        def __exit__(self, *_):
            self.close()

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

__version__ = "3.0-candump"


# ════════════════════════════════════════════════════════════════════════════
#  DBC parsing
# ════════════════════════════════════════════════════════════════════════════
BO_RE = re.compile(r"^BO_\s+(\d+)\s+(\w+)\s*:\s*(\d+)\s+(\S+)")
SG_RE = re.compile(
    r'^\s*SG_\s+(\w+)\s*:\s*(\d+)\|(\d+)@(\d)([+-])\s*'
    r'\(([^,]+),([^)]+)\)\s*\[([^|]*)\|([^\]]*)\]\s*"([^"]*)"'
)


def parse_dbc(text):
    """Return a list of message dicts with nested signal dicts."""
    messages = []
    current = None
    for raw in text.splitlines():
        m = BO_RE.match(raw)
        if m:
            mid = int(m.group(1))
            extended = mid >= 0x80000000
            if extended:
                mid -= 0x80000000
            current = dict(id=mid, name=m.group(2), dlc=int(m.group(3)),
                           sender=m.group(4), extended=extended, signals=[])
            messages.append(current)
            continue
        m = SG_RE.match(raw)
        if m and current is not None:
            current["signals"].append(dict(
                name=m.group(1), start_bit=int(m.group(2)), length=int(m.group(3)),
                little_endian=(m.group(4) == "1"), signed=(m.group(5) == "-"),
                scale=float(m.group(6)), offset=float(m.group(7)),
                min=float(m.group(8)), max=float(m.group(9)), unit=m.group(10),
            ))
    return messages


def merge_dbcs(dbc_paths, max_dbc=10):
    """Parse and merge several DBCs; first definition of an ID wins."""
    messages, seen, dups = [], set(), 0
    if len(dbc_paths) > max_dbc:
        print(f"  note: {len(dbc_paths)} DBCs given, using first {max_dbc}.")
        dbc_paths = dbc_paths[:max_dbc]
    for path in dbc_paths:
        with open(path, "r", errors="replace") as fh:
            msgs = parse_dbc(fh.read())
        added = 0
        for msg in msgs:
            if msg["id"] in seen:
                dups += 1
                continue
            seen.add(msg["id"])
            msg["source"] = os.path.basename(path)
            messages.append(msg)
            added += 1
        print(f"  {os.path.basename(path):<28} {len(msgs):>4} message(s), {added} new")
    if dups:
        print(f"  {dups} duplicate message ID(s) across databases — kept first definition")
    return messages


# ════════════════════════════════════════════════════════════════════════════
#  Log parsing  (BUSMASTER + IXXAT MiniMon), auto-detected
# ════════════════════════════════════════════════════════════════════════════
BUSMASTER_RE = re.compile(
    r"^(\d{2}):(\d{2}):(\d{2}):(\d{4})\s+(?:Rx|Tx)\s+\d+\s+"
    r"0x([0-9A-Fa-f]+)\s+([sx])\s+(\d+)((?:\s+[0-9A-Fa-f]{2})*)"
)

# BUSMASTER records the channel's configured bus speed in its header, e.g.
# '***CHANNEL 1 - PCAN-USB Driver Id 16 - 250000 bps***'. Inactive/simulated
# channels are listed too, usually as '...- 000 bps***' — skip those and
# take the first channel with a real (nonzero) speed, which in a normal
# single-channel capture is Channel 1, listed first.
CHANNEL_BAUD_RE = re.compile(r"CHANNEL\s+\d+[^\n]*?-\s*(\d+(?:\.\d+)?)\s*(k?)bps", re.I)


def declared_log_baudrate(text_head):
    """Bit/s declared in a BUSMASTER log header, or None if it's not there
    (older exports, a different tool, or no active channel found)."""
    for m in CHANNEL_BAUD_RE.finditer(text_head):
        val = float(m.group(1))
        if val <= 0:
            continue
        return val * 1000 if m.group(2) else val
    return None


CANDUMP_LOG_RE = re.compile(
    r'^\s*\(\s*(\d+(?:\.\d+)?)\s*\)\s+(\S+)\s+([0-9A-Fa-f]{3,8})(##?)([0-9A-Fa-f]*|[Rr]\d*)\s*$', re.M)
CANDUMP_LIVE_RE = re.compile(
    r'^\s*(?:\(\s*(\d+(?:\.\d+)?)\s*\)\s+)?(\S+)\s+([0-9A-Fa-f]{3,8})\s+\[(\d{1,2})\]\s*'
    r'((?:[0-9A-Fa-f]{2}\s*)*)$', re.M)


def detect_format(text_head):
    if re.search(r"IXXAT\s+MiniMon", text_head, re.I) or '"Identifier (hex)"' in text_head:
        return "minimon"
    if re.search(r"\bBUSMASTER\b", text_head, re.I):
        return "busmaster"
    if CANDUMP_LOG_RE.search(text_head) or CANDUMP_LIVE_RE.search(text_head):
        return "candump"
    # content sniff: quoted-CSV data rows => minimon
    if re.search(r'^\s*"[^"]*"\s*,\s*"[0-9A-Fa-f]+"', text_head, re.M):
        return "minimon"
    return "busmaster"


class Frame:
    __slots__ = ("t", "id", "dlc", "data", "extended")

    def __init__(self, t, cid, dlc, data, extended):
        self.t = t
        self.id = cid
        self.dlc = dlc
        self.data = data
        self.extended = extended


def parse_busmaster(text):
    frames, std, ext, bad = [], 0, 0, 0
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("***"):
            continue
        m = BUSMASTER_RE.match(line)
        if not m:
            if re.match(r"^\d{2}:\d{2}:\d{2}", line):
                bad += 1
            continue
        hh, mm, ss, frac, idhex, typ, dlc, bytestr = m.groups()
        t = int(hh) * 3600 + int(mm) * 60 + int(ss) + int(frac) / 10000.0
        cid = int(idhex, 16)
        dlc = int(dlc)
        data = bytes(int(b, 16) for b in bytestr.split()) if bytestr.strip() else b""
        if len(data) < dlc:
            bad += 1
            continue
        extended = typ == "x"
        std += not extended
        ext += extended
        frames.append(Frame(t, cid, dlc, data[:dlc], extended))
    return frames, dict(std=std, ext=ext, malformed=bad)


def _split_csv(line):
    out, cur, in_q = [], [], False
    i, n = 0, len(line)
    while i < n:
        c = line[i]
        if in_q:
            if c == '"':
                if i + 1 < n and line[i + 1] == '"':
                    cur.append('"'); i += 1
                else:
                    in_q = False
            else:
                cur.append(c)
        elif c == '"':
            in_q = True
        elif c == ",":
            out.append("".join(cur)); cur = []
        else:
            cur.append(c)
        i += 1
    out.append("".join(cur))
    return out


MINIMON_TIME_RE = re.compile(r"^\s*(\d{1,2}):(\d{2}):(\d{2})[.:](\d{1,6})\s*$")


def parse_minimon(text):
    frames, std, ext, bad = [], 0, 0, 0
    started = False
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if not started:
            if re.match(r'^"?\s*Time\s*"?\s*,', line, re.I):
                started = True
            continue
        f = _split_csv(line)
        if len(f) < 5:
            bad += 1
            continue
        tm = MINIMON_TIME_RE.match(f[0])
        if not tm:
            bad += 1
            continue
        frac = tm.group(4)
        t = int(tm.group(1)) * 3600 + int(tm.group(2)) * 60 + int(tm.group(3)) \
            + int(frac) / (10 ** len(frac))
        try:
            cid = int(f[1].strip(), 16)
        except ValueError:
            bad += 1
            continue
        datastr = f[4].strip()
        data = bytes(int(b, 16) for b in datastr.split()) if datastr else b""
        extended = "ext" in f[2].lower()
        std += not extended
        ext += extended
        frames.append(Frame(t, cid, len(data), data, extended))
    return frames, dict(std=std, ext=ext, malformed=bad)


def parse_candump(text):
    """
    Linux SocketCAN `candump` text — two sub-formats, both accepted line by
    line (a capture can even mix them):
      * `candump -l` log format   : (1699887726.123456) can0 18FEF100#0011223344556677
        Remote frames              : (1699887726.123456) can0 123#R  or  123#R4
        CAN FD frames (##)         : (1699887726.123456) can0 123##1DEADBEEF...
      * live / `-t {a,d,z}` format: [(timestamp) ]can0  123   [8]  00 11 22 33 44 55 66 77
    ID length decides Standard vs Extended, matching candump's own printing
    convention: <=3 hex digits is an 11-bit Standard ID, else 29-bit Extended
    (candump zero-pads Extended IDs out to 8 digits).
    """
    frames, std, ext, bad = [], 0, 0, 0
    seq = 0.0
    saw_any = saw_real_ts = False
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        m = CANDUMP_LOG_RE.match(line)
        if m:
            saw_any = True
            ts_s, iface, idhex, sep, payload = m.groups()
            t = float(ts_s)
            saw_real_ts = True
            extended = len(idhex) > 3
            cid = int(idhex, 16)
            if sep == "#" and payload[:1] in ("R", "r"):
                dlc = int(payload[1:]) if payload[1:].isdigit() else 0
                data = b""
            else:
                hexdata = payload[1:] if sep == "##" else payload   # ## carries a flags nibble first
                if len(hexdata) % 2:
                    bad += 1
                    continue
                data = bytes.fromhex(hexdata) if hexdata else b""
                dlc = len(data)
            std += not extended
            ext += extended
            frames.append(Frame(t, cid, dlc, data, extended))
            continue
        m = CANDUMP_LIVE_RE.match(line)
        if m:
            saw_any = True
            ts_s, iface, idhex, dlc_s, bytestr = m.groups()
            if ts_s is not None:
                t = float(ts_s)
                saw_real_ts = True
            else:
                t = seq
                seq += 0.001                       # no timestamp column — keep frame order only
            extended = len(idhex) > 3
            cid = int(idhex, 16)
            dlc = int(dlc_s)
            data = bytes(int(b, 16) for b in bytestr.split()) if bytestr.strip() else b""
            if len(data) < dlc:
                bad += 1
                continue
            std += not extended
            ext += extended
            frames.append(Frame(t, cid, dlc, data[:dlc], extended))
            continue
        bad += 1
    return frames, dict(std=std, ext=ext, malformed=bad, synthetic_time=(saw_any and not saw_real_ts))


# ── metadata (date / start / end / duration) ────────────────────────────────
def _fmt_date(raw, hint):
    parts = [p for p in re.split(r"[^0-9]+", raw) if p]
    if len(parts) < 3:
        return raw
    yi = next((i for i, x in enumerate(parts) if len(x) == 4), 2)
    year = parts[yi]
    two = [int(x) for i, x in enumerate(parts) if i != yi][:2]
    if two[0] > 12 >= two[1]:
        d, mo = two
    elif two[1] > 12 >= two[0]:
        mo, d = two
    elif hint == "mdy":
        mo, d = two
    else:
        d, mo = two
    mon = MONTHS[max(0, min(11, mo - 1))]
    return f"{d:02d} {mon} {year}"


def _to24(s):
    m = re.search(r"(\d{1,2}):(\d{2})(?::(\d{2}))?\s*(AM|PM)?", s, re.I)
    if not m:
        return s.strip()
    h = int(m.group(1))
    ap = (m.group(4) or "").upper()
    if ap == "PM" and h < 12:
        h += 12
    if ap == "AM" and h == 12:
        h = 0
    return f"{h:02d}:{m.group(2)}:{m.group(3) or '00'}"


def _sec_to_clock(t):
    t = t % 86400
    h = int(t // 3600); m = int((t % 3600) // 60); s = int(t % 60)
    ms = round((t - int(t)) * 1000)
    return f"{h:02d}:{m:02d}:{s:02d}.{ms:03d}"


def _fmt_duration(sec):
    h = int(sec // 3600); m = int((sec % 3600) // 60); s = sec % 60
    return f"{sec:.3f} s  ({h:02d}:{m:02d}:{s:06.3f})"


def parse_trim_value(val, log_t0):
    """
    Convert a trim value to seconds-from-log-start. Accepts a plain number
    (already relative seconds) or a clock string 'HH:MM:SS[.mmm]' / 'MM:SS',
    which is interpreted against the log clock (absolute minus log start).
    Returns None for blank input; raises ValueError on garbage.
    """
    if val is None:
        return None
    val = str(val).strip()
    if val == "":
        return None
    if ":" in val:
        parts = [float(p) for p in val.split(":")]
        if len(parts) == 3:
            absolute = parts[0] * 3600 + parts[1] * 60 + parts[2]
        elif len(parts) == 2:
            absolute = parts[0] * 60 + parts[1]
        else:
            absolute = parts[0]
        return absolute - log_t0
    return float(val)


FORMAT_LABELS = {"minimon": "IXXAT MiniMon", "busmaster": "BUSMASTER",
                 "candump": "candump (SocketCAN)"}


def build_meta(text, frames, fmt):
    
    meta = dict(format=FORMAT_LABELS.get(fmt, "BUSMASTER"),
                version="—", date="—", start="—", end="—",
                duration="—", dur_sec=0.0, frames=len(frames))
    first = None
    if frames:
        ts = np.fromiter((f.t for f in frames), dtype=np.float64, count=len(frames))
        first, last = float(ts.min()), float(ts.max())
        dur = last - first
        if dur < 0:
            dur += 86400
        meta["dur_sec"] = dur
        meta["duration"] = _fmt_duration(dur)
        meta["start"] = _sec_to_clock(first)
        meta["end"] = _sec_to_clock(last)
    if fmt == "minimon":
        v = re.search(r"MiniMon[^\n]*Version:\s*([\d.]+)", text, re.I)
        if v:
            meta["version"] = v.group(1)
        d = re.search(r"^\s*Date:\s*([0-9]{1,4}[\/\-.][0-9]{1,2}[\/\-.][0-9]{1,4})",
                      text, re.I | re.M)
        if d:
            meta["date"] = _fmt_date(d.group(1), "mdy")
        s = re.search(r"Start time:\s*([0-9:]+\s*(?:AM|PM)?)", text, re.I)
        e = re.search(r"Stop time:\s*([0-9:]+\s*(?:AM|PM)?)", text, re.I)
        if s:
            meta["start"] = _to24(s.group(1))
        if e:
            meta["end"] = _to24(e.group(1))
    elif fmt == "candump":
        meta["version"] = "SocketCAN candump"
        # candump timestamps are Unix epoch seconds (real calendar time), not
        # seconds-since-midnight like the other two formats — recover an
        # actual date from them when they look like a real epoch (post-2001).
        if first is not None and first > 1e9:
            gm = time.gmtime(first)
            meta["date"] = f"{gm.tm_mday:02d} {MONTHS[gm.tm_mon - 1]} {gm.tm_year}  (UTC)"
    else:
        d = re.search(r"START DATE(?:\s+AND\s+TIME)?\s*[:\-]?\s*"
                      r"([0-9]{1,2}[:\-\/][0-9]{1,2}[:\-\/][0-9]{2,4})", text, re.I)
        if d:
            meta["date"] = _fmt_date(d.group(1), "dmy")
        v = re.search(r"BUSMASTER\s+Ver\s*([0-9][0-9.]*)", text, re.I)
        if v:
            meta["version"] = v.group(1)
    return meta


def load_log(path):
    with open(path, "r", errors="replace") as fh:
        text = fh.read()
    fmt = detect_format(text[:3000])
    parser = dict(minimon=parse_minimon, candump=parse_candump).get(fmt, parse_busmaster)
    frames, counts = parser(text)
    meta = build_meta(text, frames, fmt)
    # MiniMon CSVs do carry a "Baudrate:" header line, but as an opaque
    # controller register pair (e.g. "40 2B"), not a plain bit/s value — no
    # verified table to decode that safely, so only BUSMASTER is attempted.
    meta["baudrate"] = declared_log_baudrate(text[:4000]) if fmt == "busmaster" else None
    if fmt == "candump" and counts.get("synthetic_time"):
        meta["note"] = ("No timestamps found in this candump capture — frame order was kept "
                        "but times are 1 ms apart and synthetic, so durations/rates/bus-load "
                        "below are not meaningful. Capture with 'candump -l' or 'candump -t a' "
                        "for real timestamps.")
    return fmt, frames, counts, meta


# ════════════════════════════════════════════════════════════════════════════
#  Signal decode  (vectorized over frames; validated bit-walk)
# ════════════════════════════════════════════════════════════════════════════
def _signal_bit_map(sig):
    """List of (byte_idx, bit_in_byte, dest_shift) — exact Vector layout."""
    length = sig["length"]
    out = []
    if sig["little_endian"]:                          # Intel: LSB-first, linear
        for i in range(length):
            N = sig["start_bit"] + i
            out.append((N >> 3, N & 7, i))
    else:                                             # Motorola: MSB-first sawtooth
        bit_pos = sig["start_bit"]
        for i in range(length):
            out.append((bit_pos >> 3, bit_pos & 7, length - 1 - i))
            bit_pos = bit_pos + 15 if (bit_pos & 7) == 0 else bit_pos - 1
    return out


def decode_signal_vec(data_matrix, sig):
    """
    data_matrix : uint8 array (n_frames, width) already zero-padded to the
                  message DLC.  Returns float64 array of engineering values.
    """
    length = sig["length"]
    n, width = data_matrix.shape
    if length > 63:                                   # rare wide signal -> Python big-int path
        return _decode_signal_slow(data_matrix, sig)
    raw = np.zeros(n, dtype=np.int64)
    for byte_idx, bit_in_byte, dest in _signal_bit_map(sig):
        if byte_idx >= width:
            continue
        bit = (data_matrix[:, byte_idx] >> bit_in_byte) & 1
        raw |= bit.astype(np.int64) << dest
    if sig["signed"]:
        sign = np.int64(1) << (length - 1)
        raw = np.where((raw & sign) != 0, raw - (np.int64(1) << length), raw)
    return raw.astype(np.float64) * sig["scale"] + sig["offset"]


def _decode_signal_slow(data_matrix, sig):
    length = sig["length"]
    bmap = _signal_bit_map(sig)
    out = np.empty(data_matrix.shape[0], dtype=np.float64)
    for r in range(data_matrix.shape[0]):
        row = data_matrix[r]
        raw = 0
        for byte_idx, bit_in_byte, dest in bmap:
            if byte_idx < len(row):
                raw |= int((row[byte_idx] >> bit_in_byte) & 1) << dest
        if sig["signed"] and (raw & (1 << (length - 1))):
            raw -= (1 << length)
        out[r] = raw * sig["scale"] + sig["offset"]
    return out


def decode_all(frames, messages, wanted=None):
    """
    Group frames by ID, then decode every (selected) signal on each matched
    message.  Returns list of signal series dicts:
        {name, unit, msg_id, t (ndarray), v (ndarray)}
    """
    frame_ids = set(f.id for f in frames)
    matched_msgs = [m for m in messages if m["id"] in frame_ids]

    # bucket frame indices by id
    buckets = {}
    for i, f in enumerate(frames):
        buckets.setdefault(f.id, []).append(i)

    series = []
    for msg in tqdm(matched_msgs, desc="Decoding", unit="msg"):
        idxs = buckets.get(msg["id"], [])
        if not idxs:
            continue
        width = msg["dlc"]
        # build (n, width) uint8 matrix, zero-padded
        mat = np.zeros((len(idxs), width), dtype=np.uint8)
        t = np.empty(len(idxs), dtype=np.float64)
        for r, fi in enumerate(idxs):
            fr = frames[fi]
            t[r] = fr.t
            b = fr.data[:width]
            mat[r, :len(b)] = np.frombuffer(b, dtype=np.uint8)
        order = np.argsort(t, kind="stable")
        t = t[order]
        mat = mat[order]
        for sig in msg["signals"]:
            if wanted and sig["name"] not in wanted:
                continue
            v = decode_signal_vec(mat, sig)
            series.append(dict(name=sig["name"], unit=sig["unit"],
                               msg_id=msg["id"], t=t, v=v))
    return series, matched_msgs


# ════════════════════════════════════════════════════════════════════════════
#  Stats + merge
# ════════════════════════════════════════════════════════════════════════════
def stats(v):
    return dict(n=len(v), min=float(v.min()), max=float(v.max()),
                mean=float(v.mean()), std=float(v.std()))


def merge_asof_nearest(times, stimes, svals, tol):
    """Nearest-value join within tolerance; NaN where nothing is close."""
    if len(stimes) == 0:
        return np.full(len(times), np.nan)
    idx = np.clip(np.searchsorted(stimes, times), 0, len(stimes) - 1)
    left = np.clip(idx - 1, 0, len(stimes) - 1)
    d_right = np.abs(stimes[idx] - times)
    d_left = np.abs(stimes[left] - times)
    use_left = d_left < d_right
    best = np.where(use_left, left, idx)
    bestd = np.where(use_left, d_left, d_right)
    return np.where(bestd <= tol, svals[best], np.nan)


# ════════════════════════════════════════════════════════════════════════════
#  Bus health  (per-ID timing is exact; bus load % is an estimate — see
#  BUS_LOAD_FORMULA)
# ════════════════════════════════════════════════════════════════════════════
BUS_LOAD_FORMULA = (
    "Bus load % = (bits estimated in each 1-second window) / "
    "(assumed bit rate x 1 s) x 100.\n"
    "Per frame: (47 overhead bits for a standard frame, 67 for extended — "
    "SOF + arbitration + control + CRC + ACK + EOF + IFS, per ISO 11898-1) "
    "+ 8 x DLC data bits, then x1.1 for average bit-stuffing inflation "
    "(a plain ASCII trace doesn't record each frame's true stuffed length).\n"
    "This is an estimate for spotting load trends and hotspots, not a "
    "certified measurement — accuracy depends entirely on the assumed bit "
    "rate matching the real bus."
)
_STUFF_FACTOR = 1.1
_GAP_THRESHOLD_DEFAULT = 2.0


def detect_bus_off_gaps(frames, threshold_sec=_GAP_THRESHOLD_DEFAULT):
    """
    Silent windows across the WHOLE trace (no frame of ANY ID) longer than
    threshold_sec — the signature of a bus going off, or the logger/vehicle
    losing power mid-capture, as opposed to a single ID dropping out while
    the rest of the bus keeps talking (see the per-ID "dropout" flag for that).

    Returns a list of dicts sorted by time: {start_t, end_t, duration,
    start_clock, end_clock} — start_t/end_t are the timestamps of the last
    frame before, and first frame after, the silence (same units as Frame.t).
    """
    if len(frames) < 2:
        return []
    ts = np.sort(np.fromiter((f.t for f in frames), float, len(frames)))
    gaps = np.diff(ts)
    events = []
    for i in np.nonzero(gaps > threshold_sec)[0]:
        start_t, end_t = float(ts[i]), float(ts[i + 1])
        events.append(dict(start_t=start_t, end_t=end_t, duration=end_t - start_t,
                           start_clock=_sec_to_clock(start_t), end_clock=_sec_to_clock(end_t)))
    return events


def analyze_bus_health(frames, dur_sec, frame_name_by_id, baudrate=500000.0,
                       stuff_factor=_STUFF_FACTOR, window_sec=1.0, dropout_mult=3.0,
                       gap_threshold_sec=_GAP_THRESHOLD_DEFAULT):
    """
    Per-ID frame timing (exact, from frame timestamps) plus an estimated
    bus-load-over-time (see BUS_LOAD_FORMULA for exactly how), plus
    whole-bus silence gaps (see detect_bus_off_gaps).

    Returns dict(per_id, bins, load, load_avg, load_peak, bus_off,
                 gap_threshold_sec, bins_t0):
      per_id : list of dicts sorted by frame count desc —
               {id, name, count, hz, avg_gap_ms, std_gap_ms, max_gap_ms, dropout}
               A "dropout" flag means that ID's longest gap was more than
               dropout_mult x its own median gap — excluding any gap that
               overlaps a BUS-OFF/Power-OFF event, since that's the whole
               bus going quiet, not this ID uniquely misbehaving. max_gap_ms
               itself still reports the true longest gap either way.
      bins, load : ndarrays, one point per window_sec bin — bin start (s
               from the first frame) and estimated load (%).
      load_avg, load_peak : float, percent.
      bus_off : list of gaps (see detect_bus_off_gaps) — flag these as
               BUS-OFF / Power-OFF in a report; bin/gap times share the same
               origin (bins_t0), so a gap at start_t-bins_t0 lines up with
               the bins x-axis directly.
    """
    if not frames or dur_sec <= 0:
        return dict(per_id=[], bins=np.array([]), load=np.array([]),
                    load_avg=0.0, load_peak=0.0, bus_off=[],
                    gap_threshold_sec=gap_threshold_sec, bins_t0=0.0)

    ts = np.fromiter((f.t for f in frames), float, len(frames))
    t0 = float(ts.min())

    # computed up front so the per-ID dropout check below can exclude gaps
    # that are just that ID going quiet because the WHOLE bus went quiet
    # (a BUS-OFF / Power-OFF event) — otherwise one bus-wide outage would
    # mark every ID active around it as individually "dropped out"
    bus_off = detect_bus_off_gaps(frames, gap_threshold_sec)

    def _max_gap_excluding_busoff(times, gaps):
        if not bus_off or not len(gaps):
            return float(gaps.max()) if len(gaps) else 0.0
        starts, ends = times[:-1], times[1:]
        overlap = np.zeros(len(gaps), dtype=bool)
        for g in bus_off:
            overlap |= (starts < g["end_t"]) & (ends > g["start_t"])
        remaining = gaps[~overlap]
        return float(remaining.max()) if len(remaining) else 0.0

    by_id = {}
    for f in frames:
        by_id.setdefault(f.id, []).append(f.t)
    per_id = []
    for cid, times in by_id.items():
        times = np.sort(np.asarray(times))
        n = len(times)
        gaps = np.diff(times) * 1000.0                      # ms
        dropout = False
        if len(gaps) >= 4:
            med = float(np.median(gaps))
            max_own_gap = _max_gap_excluding_busoff(times, gaps)
            if med > 0 and max_own_gap > dropout_mult * med:
                dropout = True
        per_id.append(dict(
            id=cid, name=frame_name_by_id.get(cid, ""), count=n,
            hz=n / dur_sec,
            avg_gap_ms=float(gaps.mean()) if len(gaps) else 0.0,
            std_gap_ms=float(gaps.std()) if len(gaps) else 0.0,
            max_gap_ms=float(gaps.max()) if len(gaps) else 0.0,
            dropout=dropout,
        ))
    per_id.sort(key=lambda r: -r["count"])

    def _frame_bits(f):
        base = 67 if f.extended else 47
        return (base + 8 * f.dlc) * stuff_factor

    n_bins = max(1, int(np.ceil(dur_sec / window_sec)))
    bin_bits = np.zeros(n_bins)
    for f in frames:
        idx = min(n_bins - 1, max(0, int((f.t - t0) / window_sec)))
        bin_bits[idx] += _frame_bits(f)
    # the last bin usually covers less than a full window_sec of real time
    # (dur_sec rarely divides evenly) — size it to its actual span, or it's
    # systematically undercounted against a full window's capacity
    bin_durs = np.full(n_bins, window_sec)
    bin_durs[-1] = dur_sec - (n_bins - 1) * window_sec
    load_pct = bin_bits / (bin_durs * baudrate) * 100.0
    bin_starts = np.arange(n_bins) * window_sec

    return dict(per_id=per_id, bins=bin_starts, load=load_pct,
               load_avg=float(load_pct.mean()), load_peak=float(load_pct.max()),
               bus_off=bus_off, gap_threshold_sec=gap_threshold_sec, bins_t0=t0)


# ════════════════════════════════════════════════════════════════════════════
#  Plots  ("progress graphs")
# ════════════════════════════════════════════════════════════════════════════
PALETTE = ["#d98a00", "#0e9bb8", "#b5259e", "#2aa14a",
           "#1f77d8", "#e8433d", "#7b5cff", "#8a8f00"]


def _decimate(t, v, max_n=6000):
    if len(t) <= max_n:
        return t, v
    step = int(np.ceil(len(t) / max_n))
    return t[::step], v[::step]


def make_plots(series, outdir, t0, show=False):
    try:
        import matplotlib
        if not show:
            matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        print("  matplotlib not installed — skipping plots (pip install matplotlib)")
        return []

    plotdir = os.path.join(outdir, "plots")
    os.makedirs(plotdir, exist_ok=True)
    paths = []

    # one graph per signal
    for s in tqdm(series, desc="Plotting", unit="sig"):
        col = PALETTE[len(paths) % len(PALETTE)]
        t, v = _decimate(s["t"] - t0, s["v"])
        fig, ax = plt.subplots(figsize=(9, 3.2), dpi=110)
        ax.plot(t, v, color=col, lw=1.1)
        st = stats(s["v"])
        title = s["name"] + (f"  [{s['unit']}]" if s["unit"] else "")
        ax.set_title(title, fontsize=11, fontweight="bold")
        ax.set_xlabel("t (s)", fontsize=9)
        ax.set_ylabel(s["unit"] or "value", fontsize=9)
        ax.grid(True, alpha=0.25)
        ax.text(0.99, 0.02,
                f"min {st['min']:.3f}   max {st['max']:.3f}   avg {st['mean']:.3f}",
                transform=ax.transAxes, ha="right", va="bottom",
                fontsize=8, color="#555")
        fig.tight_layout()
        safe = re.sub(r"[^\w.-]", "_", s["name"])
        p = os.path.join(plotdir, f"{safe}.png")
        fig.savefig(p)
        plt.close(fig)
        paths.append(p)

    # overlay (independent y-axes, up to first 6 signals) + normalized overlay
    overlay = series[:6]
    if len(overlay) >= 2:
        # raw overlay with twin axes
        fig, ax0 = plt.subplots(figsize=(11, 4.2), dpi=110)
        axes = [ax0]
        for i, s in enumerate(overlay):
            ax = ax0 if i == 0 else ax0.twinx()
            if i >= 2:
                ax.spines["right"].set_position(("outward", 46 * (i - 1)))
            col = PALETTE[i % len(PALETTE)]
            t, v = _decimate(s["t"] - t0, s["v"])
            ax.plot(t, v, color=col, lw=1.1,
                    label=s["name"] + (f" [{s['unit']}]" if s["unit"] else ""))
            ax.set_ylabel(s["unit"] or s["name"], color=col, fontsize=9)
            ax.tick_params(axis="y", labelcolor=col)
            axes.append(ax)
        ax0.set_xlabel("t (s)")
        ax0.set_title("Signal overlay", fontsize=11, fontweight="bold")
        ax0.grid(True, alpha=0.25)
        lines = [ln for ax in axes for ln in ax.get_lines()]
        ax0.legend(lines, [ln.get_label() for ln in lines], fontsize=8, loc="upper right")
        fig.tight_layout()
        p = os.path.join(plotdir, "_overlay.png")
        fig.savefig(p); plt.close(fig); paths.append(p)

        # normalized overlay (shape comparison)
        fig, ax = plt.subplots(figsize=(11, 4.2), dpi=110)
        for i, s in enumerate(overlay):
            col = PALETTE[i % len(PALETTE)]
            t, v = _decimate(s["t"] - t0, s["v"])
            rng = (v.max() - v.min()) or 1.0
            ax.plot(t, (v - v.min()) / rng, color=col, lw=1.1,
                    label=s["name"])
        ax.set_xlabel("t (s)"); ax.set_ylabel("normalized 0–1")
        ax.set_title("Signal overlay (normalized)", fontsize=11, fontweight="bold")
        ax.grid(True, alpha=0.25); ax.legend(fontsize=8, loc="upper right")
        fig.tight_layout()
        p = os.path.join(plotdir, "_overlay_norm.png")
        fig.savefig(p); plt.close(fig); paths.append(p)

    if show:
        import matplotlib.pyplot as plt
        # re-show a lightweight overlay interactively
        if len(overlay) >= 2:
            plt.figure(figsize=(11, 4.5))
            for i, s in enumerate(overlay):
                t, v = _decimate(s["t"] - t0, s["v"])
                rng = (v.max() - v.min()) or 1.0
                plt.plot(t, (v - v.min()) / rng, lw=1.1, label=s["name"],
                         color=PALETTE[i % len(PALETTE)])
            plt.legend(); plt.xlabel("t (s)"); plt.ylabel("normalized")
            plt.title("Signal overlay (normalized)"); plt.grid(alpha=0.25)
            plt.tight_layout(); plt.show()
    return paths


# ════════════════════════════════════════════════════════════════════════════
#  Excel export
# ════════════════════════════════════════════════════════════════════════════
def _safe_sheet(name, used):
    n = re.sub(r"[:\\/?*\[\]]", " ", name).strip()[:31] or "Signal"
    out, i = n, 2
    while out.lower() in used:
        out = (n[:28] + "_" + str(i))[:31]
        i += 1
    used.add(out.lower())
    return out


def export_excel(path, meta, dbc_names, series, frames, frame_name_by_id,
                 tol, plot_paths, t0, sections=None, dynamics=None, faults=None,
                 progress=None, components=None, dyn_chart=None, bus_health=None):
    """
    sections : optional set of {'summary','frames','merged',
               'charts','dynamics','faults','components','bus_health'}
               controlling which sheets are written (Log Info is always
               included). The Merged sheet carries every signal in `series`
               (typically the full decoded DBC set); there are no more
               per-signal sheets.
    dynamics : optional list of (metric, value, unit) rows -> 'Vehicle Dynamics'.
    dyn_chart : optional path to a PNG (power / cumulative energy over time)
               embedded alongside the Vehicle Dynamics table.
    faults   : optional list of fault-episode rows -> 'Faults'.
    components : optional list of (component, signal, unit, n, min, max, mean,
               std) rows -> 'Components' (written when 'components' in sections).
    bus_health : optional dict as returned by analyze_bus_health(), plus a
               'baudrate' key (the bit rate it was computed with) and an
               optional 'chart' key (path to a bus-load-over-time PNG) ->
               'Bus Health' (written when 'bus_health' in sections).
    progress : optional callable(fraction: float, label: str) for UI progress.
    meta may also carry 'vin', 'payload', 'description' for the header.
    """
    def _p(frac, label=""):
        if progress:
            try:
                progress(frac, label)
            except Exception:
                pass

    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Font
    except ImportError:
        print("  openpyxl not installed — skipping Excel (pip install openpyxl)")
        return None

    if sections is None:
        sections = {"summary", "frames", "merged", "charts"}
    bold = Font(bold=True)
    _p(0.02, "Writing header…")

    wb = Workbook()
    used = set()

    # Log Info (always)
    ws = wb.active
    ws.title = _safe_sheet("Log Info", used)
    info = []
    if meta.get("vin"):
        info.append(("VIN", meta["vin"]))
    if meta.get("payload") not in (None, ""):
        info.append(("Payload (load body)", meta["payload"]))
    info += [
        ("Format", meta["format"]), ("Version", meta["version"]),
        ("Date", meta["date"]), ("Start time", meta["start"]),
        ("End time", meta["end"]), ("Duration", meta["duration"]),
        ("Trim window", meta.get("trim", "full log")),
        ("Total frames", len(frames)),
        ("DBC databases", ", ".join(dbc_names)),
        ("Message IDs matched",
         ", ".join(sorted({f"0x{s['msg_id']:X}" for s in series}))),
        ("Signals decoded", len(series)),
    ]
    if meta.get("description"):
        info.append(("Description", meta["description"]))
    for k, v in info:
        ws.append([k, v])
    for row in ws.iter_rows(min_col=1, max_col=1):
        row[0].font = bold
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 46

    # Bus Health (per-ID timing is exact; bus load % is an estimate)
    if bus_health and "bus_health" in sections:
        ws = wb.create_sheet(_safe_sheet("Bus Health", used))
        baudrate = bus_health.get("baudrate", 500000.0)
        bus_off = bus_health.get("bus_off") or []
        gap_threshold = bus_health.get("gap_threshold_sec", _GAP_THRESHOLD_DEFAULT)
        overload = bus_health["load_peak"] > 100.0
        if overload:
            ws.append([f"⚠ Estimated peak bus load is {bus_health['load_peak']:.0f}%, which is "
                       f"impossible on a real bus — the assumed bit rate is very likely too low."])
            ws[f"A{ws.max_row}"].font = Font(bold=True, color="CC0000")
        if bus_off:
            ws.append([f"⚠ {len(bus_off)} BUS-OFF / Power-OFF gap(s) detected — no frames of any "
                       f"ID for longer than {gap_threshold:g}s. See the table below for exact times."])
            ws[f"A{ws.max_row}"].font = Font(bold=True, color="CC0000")
        summary_rows = [
            ("Bit rate assumed", f"{baudrate/1000:.0f} kbit/s"),
            ("Average bus load (est.)", f"{bus_health['load_avg']:.2f} %"),
            ("Peak bus load (est.)", f"{bus_health['load_peak']:.2f} %"),
            ("Unique message IDs seen", len(bus_health["per_id"])),
            ("BUS-OFF / Power-OFF gaps", len(bus_off)),
        ]
        for label, value in summary_rows:
            ws.append([label, value])
            ws[f"A{ws.max_row}"].font = bold
        ws.append([])
        ws.append(["Bus load % — how it's calculated"])
        ws[f"A{ws.max_row}"].font = bold
        for line in BUS_LOAD_FORMULA.split("\n"):
            ws.append([line])
        ws.append([])
        ws.append(["Per-Message-ID Timing"])
        ws[f"A{ws.max_row}"].font = bold
        ws.append(["CAN_ID", "Message", "Count", "Freq (Hz)", "Avg gap (ms)",
                   "Std gap (ms)", "Max gap (ms)", "Dropout?"])
        for cell in ws[ws.max_row]:
            cell.font = bold
        for r in bus_health["per_id"]:
            ws.append([f"0x{r['id']:X}", r["name"], r["count"],
                       round(r["hz"], 3), round(r["avg_gap_ms"], 2),
                       round(r["std_gap_ms"], 2), round(r["max_gap_ms"], 2),
                       "YES" if r["dropout"] else ""])
        for col, w in zip("ABCDEFGH", (12, 24, 10, 11, 13, 13, 13, 10)):
            ws.column_dimensions[col].width = w
        ws.append([])
        ws.append([f"BUS-OFF / Power-OFF Events  (no frames of any ID for > {gap_threshold:g} s)"])
        ws[f"A{ws.max_row}"].font = bold
        if bus_off:
            ws.append(["Start", "End", "Duration (s)", "Start t (s)", "End t (s)"])
            for cell in ws[ws.max_row]:
                cell.font = bold
            for g in bus_off:
                ws.append([g["start_clock"], g["end_clock"], round(g["duration"], 3),
                           round(g["start_t"], 3), round(g["end_t"], 3)])
        else:
            ws.append(["None detected — frames were present continuously throughout the trace."])
        bus_chart = bus_health.get("chart")
        if bus_chart and os.path.exists(bus_chart):
            try:
                img = XLImage(bus_chart)
                img.width, img.height = 680, 260
                ws.add_image(img, "J2")
            except Exception:
                pass

    # Vehicle Dynamics
    if dynamics and "dynamics" in sections:
        ws = wb.create_sheet(_safe_sheet("Vehicle Dynamics", used))
        ws.append(["Metric", "Value", "Unit"])
        for cell in ws[1]:
            cell.font = bold
        for metric, value, unit in dynamics:
            ws.append([metric, value, unit])
        ws.column_dimensions["A"].width = 32
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 10
        if dyn_chart and os.path.exists(dyn_chart):
            try:
                img = XLImage(dyn_chart)
                img.width, img.height = 680, 300
                ws.add_image(img, "E2")
            except Exception:
                pass

    # Components (per-subsystem signal summary)
    if components and "components" in sections:
        ws = wb.create_sheet(_safe_sheet("Components", used))
        ws.append(["Component", "Signal", "Unit", "Samples",
                   "Min", "Max", "Mean", "Std Dev"])
        for cell in ws[1]:
            cell.font = bold
        for r in components:
            ws.append(list(r))
        for col, w in zip("ABCDEFGH", (14, 26, 10, 10, 12, 12, 12, 12)):
            ws.column_dimensions[col].width = w

    # Faults (accumulated fault episodes with start / end time)
    if faults and "faults" in sections:
        ws = wb.create_sheet(_safe_sheet("Faults", used))
        ws.append(["Signal", "Code / value", "Start", "End",
                   "Duration (s)", "Start t (s)", "End t (s)"])
        for cell in ws[1]:
            cell.font = bold
        for r in faults:
            ws.append(list(r))
        for col, w in zip("ABCDEFG", (26, 14, 15, 15, 13, 12, 12)):
            ws.column_dimensions[col].width = w

    # Summary
    if "summary" in sections:
        ws = wb.create_sheet(_safe_sheet("Summary", used))
        ws.append(["Signal", "Unit", "Samples", "Min", "Max", "Mean", "Std Dev"])
        for cell in ws[1]:
            cell.font = bold
        for s in series:
            st = stats(s["v"])
            ws.append([s["name"], s["unit"], st["n"],
                       round(st["min"], 4), round(st["max"], 4),
                       round(st["mean"], 4), round(st["std"], 4)])

    # CAN Frames (full raw trace) — Excel hard-caps a sheet at 1,048,576 rows
    # (including the header); a trace with more frames than that would lose
    # data — or produce an invalid file — on a single sheet. Spill into
    # "CAN Frames_2", "CAN Frames_3", ... so every frame still lands somewhere.
    if "frames" in sections:
        _p(0.06, "Writing CAN frames…")
        FRAME_HEADER = ["Time_s", "Clock", "CAN_ID", "Type", "Message", "DLC",
                        "Data (hex)", "B0", "B1", "B2", "B3", "B4", "B5", "B6", "B7"]
        FRAMES_PER_SHEET = 1_000_000  # headroom under Excel's 1,048,576 row cap
        ft0 = t0
        nf = max(1, len(frames))
        ws = None
        for idx, f in enumerate(tqdm(frames, desc="Excel: frames", unit="fr")):
            if idx % FRAMES_PER_SHEET == 0:
                ws = wb.create_sheet(_safe_sheet("CAN Frames", used))
                ws.append(FRAME_HEADER)
                for cell in ws[1]:
                    cell.font = bold
            hexbytes = [f"{b:02X}" for b in f.data]
            row = [round(f.t - ft0, 4), _sec_to_clock(f.t), f"0x{f.id:X}",
                   "Ext" if f.extended else "Std",
                   frame_name_by_id.get(f.id, ""), f.dlc, " ".join(hexbytes)]
            row += hexbytes[:8] + [""] * (8 - len(hexbytes[:8]))
            ws.append(row)
            if progress and idx % 3000 == 0:
                _p(0.06 + 0.44 * idx / nf, "Writing CAN frames…")
        if ws is None:                          # no frames at all — keep the header-only sheet
            ws = wb.create_sheet(_safe_sheet("CAN Frames", used))
            ws.append(FRAME_HEADER)
            for cell in ws[1]:
                cell.font = bold
        frame_sheet_count = max(1, -(-len(frames) // FRAMES_PER_SHEET))
        if frame_sheet_count > 1:
            print(f"  CAN Frames split across {frame_sheet_count} sheets "
                  f"({len(frames):,} frames, {FRAMES_PER_SHEET:,} per sheet — "
                  f"Excel's own limit is 1,048,576 rows/sheet).")

    # Merged (as-of, nearest within tol) — split across sheets instead of
    # skipping when it's too big: by signal (column) groups to stay under a
    # sane per-sheet cell budget, and by time (row) groups too if the trace
    # has more unique timestamps than Excel allows on one sheet (1,048,576
    # rows) — so nothing is ever dropped, it just spills onto more sheets.
    if series and "merged" in sections:
        _p(0.52, "Writing merged sheet…")
        all_times = np.unique(np.round(
            np.concatenate([s["t"] - t0 for s in series]), 4))
        MERGE_CELL_BUDGET = 1_500_000
        MERGE_ROW_CAP = 1_000_000        # headroom under Excel's 1,048,576 row limit

        cols = []
        for s in tqdm(series, desc="Excel: merge", unit="sig"):
            cols.append(merge_asof_nearest(all_times, s["t"] - t0, s["v"], tol))
        cols = np.array(cols)                            # (nsig, ntimes)

        n_times = len(all_times)
        row_groups = [(i, min(i + MERGE_ROW_CAP, n_times))
                     for i in range(0, n_times, MERGE_ROW_CAP)] or [(0, 0)]
        merge_plan = []
        for row_start, row_end in row_groups:
            signals_per_sheet = max(1, MERGE_CELL_BUDGET // max(1, row_end - row_start))
            for col_start in range(0, len(series), signals_per_sheet):
                merge_plan.append((row_start, row_end, col_start,
                                   min(col_start + signals_per_sheet, len(series))))

        nt = max(1, n_times)
        for pi, (row_start, row_end, col_start, col_end) in enumerate(merge_plan):
            if len(merge_plan) == 1:
                label = "Merged"
            elif len(row_groups) > 1:
                label = f"Merged t{row_start+1}-{row_end} sig{col_start+1}-{col_end}"
            else:
                label = f"Merged signals {col_start+1}-{col_end}"
            ws = wb.create_sheet(_safe_sheet(label, used))
            ws.append(["Time_s"] + [f"{s['name']} [{s['unit']}]" if s["unit"] else s["name"]
                                    for s in series[col_start:col_end]])
            for cell in ws[1]:
                cell.font = bold
            for i in range(row_start, row_end):
                row = [round(float(all_times[i]), 4)]
                for c in range(col_start, col_end):
                    val = cols[c, i]
                    row.append("" if np.isnan(val) else round(float(val), 4))
                ws.append(row)
                if progress and i % 4000 == 0:
                    _p(0.55 + 0.15 * i / nt, f"Writing merged sheet {pi+1}/{len(merge_plan)}…")
        if len(merge_plan) > 1:
            print(f"  Merged data split across {len(merge_plan)} sheets "
                  f"({n_times:,} time points x {len(series)} signals — "
                  f"Excel's own limit is 1,048,576 rows/sheet).")

    # Charts sheet — embed the per-signal PNGs
    single = [p for p in plot_paths if not os.path.basename(p).startswith("_")]
    if single and "charts" in sections:
        _p(0.92, "Embedding charts…")
        ws = wb.create_sheet(_safe_sheet("Charts", used))
        r = 1
        for p in single:
            try:
                img = XLImage(p)
                img.width, img.height = 640, 228
                ws.add_image(img, f"A{r}")
                r += 13
            except Exception:
                pass

    _p(0.97, "Saving workbook…")
    wb.save(path)
    _p(1.0, "")
    return path


# ════════════════════════════════════════════════════════════════════════════
#  Main
# ════════════════════════════════════════════════════════════════════════════
def interactive_setup(args):
    """
    Fill in any missing --dbc / --log / --out when the script is launched
    without them (e.g. double-clicked or run with no arguments). Tries native
    file-picker dialogs first, then falls back to typed paths.
    """
    print("No files given on the command line — entering interactive mode.\n"
          "  (tip: you can also run it as:  python can_log_analyzer.py "
          "--dbc a.dbc --log trace.csv --out ./out)\n")

    dbc, log, out = args.dbc, args.log, args.out
    used_dialog = False
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        used_dialog = True

        if not dbc:
            picked = filedialog.askopenfilenames(
                title="Select DBC database file(s) — up to 10",
                filetypes=[("DBC files", "*.dbc"), ("All files", "*.*")])
            dbc = list(picked)
        if not log:
            log = filedialog.askopenfilename(
                title="Select CAN log  (BUSMASTER .log/.asc  or  MiniMon .csv)",
                filetypes=[("CAN logs", "*.log *.asc *.csv *.txt"), ("All files", "*.*")])
        if not out:
            chosen = filedialog.askdirectory(
                title="Select output folder  (Cancel = ./can_out next to the log)")
            out = chosen or None
        root.destroy()
    except Exception:
        used_dialog = False  # headless / no Tk — fall through to text prompts

    if not used_dialog:
        if not dbc:
            raw = input("Path(s) to DBC file(s) (comma-separated): ").strip()
            dbc = [p.strip().strip('"') for p in raw.split(",") if p.strip()]
        if not log:
            log = input("Path to CAN log file: ").strip().strip('"')
        if not out:
            o = input("Output folder [./can_out]: ").strip().strip('"')
            out = o or None

    # default output next to the log if still unset
    if not out:
        out = os.path.join(os.path.dirname(os.path.abspath(log)) if log else ".", "can_out")

    # optional trim / signal filter, only prompted interactively
    if args.start is None and args.end is None:
        try:
            print("\nOptional trim — press Enter to skip.")
            print("  Enter seconds from start (e.g. 5) or clock time (e.g. 15:27:00).")
            ts = input("  Trim start: ").strip()
            te = input("  Trim end  : ").strip()
            if ts:
                args.start = ts
            if te:
                args.end = te
        except EOFError:
            pass

    args.dbc, args.log, args.out = dbc, log, out
    return args


def main():
    ap = argparse.ArgumentParser(
        description="DBC-guided CAN log decoder (BUSMASTER + IXXAT MiniMon).",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    ap.add_argument("--dbc", nargs="+", default=None, help="one or more .dbc files (max 10)")
    ap.add_argument("--log", default=None, help="BUSMASTER .log/.asc or MiniMon .csv")
    ap.add_argument("--out", default=None, help="output directory (default ./can_out)")
    ap.add_argument("--signals", nargs="*", default=None,
                    help="only decode these signal names (default: all)")
    ap.add_argument("--tol", type=float, default=0.1,
                    help="merge tolerance in seconds for the Merged sheet")
    ap.add_argument("--start", default=None,
                    help="trim start — seconds from log start (e.g. 5) or clock time (e.g. 15:27:00)")
    ap.add_argument("--end", default=None,
                    help="trim end — seconds from log start (e.g. 30) or clock time (e.g. 15:28:30)")
    ap.add_argument("--no-excel", action="store_true", help="skip the .xlsx export")
    ap.add_argument("--no-plots", action="store_true", help="skip PNG plots")
    ap.add_argument("--show", action="store_true", help="also show an overlay interactively")
    ap.add_argument("--bus-health", action="store_true",
                    help="add a Bus Health sheet: per-ID timing, bus load %%, "
                         "and BUS-OFF/Power-OFF gaps (silence across all IDs)")
    ap.add_argument("--baudrate", type=float, default=500000.0,
                    help="assumed CAN bus bit rate (bit/s), used for the bus-load estimate "
                         "(overridden by the log header when BUSMASTER recorded one)")
    ap.add_argument("--gap-threshold", type=float, default=2.0,
                    help="seconds with no frame of ANY ID before flagging BUS-OFF/Power-OFF")
    ap.add_argument("--version", action="version",
                    version=f"can_log_analyzer {__version__}")
    args = ap.parse_args()

    print(f"CAN Log Analyzer  v{__version__}")
    print("=" * 46)

    interactive = not (args.dbc and args.log)
    if interactive:
        args = interactive_setup(args)

    if not args.dbc:
        sys.exit("No DBC file selected — nothing to decode.")
    if not args.log:
        sys.exit("No log file selected — nothing to decode.")
    if args.out is None:
        args.out = "./can_out"

    t_start = time.time()
    os.makedirs(args.out, exist_ok=True)

    print("Parsing DBC database(s)…")
    messages = merge_dbcs(args.dbc)
    total_sigs = sum(len(m["signals"]) for m in messages)
    print(f"  merged: {len(messages)} unique message(s), {total_sigs} signal(s)\n")

    print(f"Loading log: {args.log}")
    fmt, frames, counts, meta = load_log(args.log)
    fmt_label = {"minimon": "IXXAT MiniMon CSV", "candump": "candump (SocketCAN)"}.get(fmt, "BUSMASTER")
    print(f"  detected format : {fmt_label}")
    if not frames:
        sys.exit(
            "  no frames parsed — this file doesn't look like a supported trace format.\n"
            "  Supported formats: BUSMASTER .log, MiniMon-compatible quoted CSV, and candump/SocketCAN log.\n"
            "  Need another format supported? Email dhanoosh2001@gmail.com with your email ID so we can update you."
        )
    print(f"  validation      : {len(frames)} frame(s) — "
          f"{counts['std']} std, {counts['ext']} ext, {counts['malformed']} skipped")
    print(f"  date/duration   : {meta['date']}  |  {meta['duration']}")
    print(f"  window          : {meta['start']} -> {meta['end']}\n")
    if meta.get("note"):
        print(f"  note: {meta['note']}\n")

    # ── optional time trim (seconds from start, or clock time) ──────────────
    log_t0 = min(f.t for f in frames)
    trim_desc = "full log"
    if args.start is not None or args.end is not None:
        try:
            rs = parse_trim_value(args.start, log_t0)
            re_ = parse_trim_value(args.end, log_t0)
        except ValueError:
            sys.exit("  invalid trim value — use seconds (e.g. 5) or clock time "
                     "(e.g. 15:27:00).")
        a = 0.0 if rs is None else max(0.0, rs)
        b = meta["dur_sec"] if re_ is None else min(meta["dur_sec"], re_)
        if b <= a:
            sys.exit(f"  invalid trim: end ({b:.3f}s) must be after start ({a:.3f}s).")
        before = len(frames)
        frames = [f for f in frames if a <= (f.t - log_t0) <= b]
        if not frames:
            sys.exit(f"  trim window {a:.3f}-{b:.3f}s contains no frames.")
        trim_desc = (f"{a:.3f}-{b:.3f} s  "
                     f"({_sec_to_clock(log_t0 + a)}–{_sec_to_clock(log_t0 + b)})")
        print(f"Trim window       : {trim_desc}  "
              f"({len(frames)} of {before} frames kept)\n")
    meta["trim"] = trim_desc

    wanted = set(args.signals) if args.signals else None
    series, matched = decode_all(frames, messages, wanted)
    if not series:
        sys.exit("No signals decoded — no DBC message IDs present in the log"
                 + (" / trim window." if trim_desc != "full log" else "."))
    print(f"  matched {len(matched)} message ID(s), decoded {len(series)} signal(s)\n")

    frame_name_by_id = {m["id"]: m["name"] for m in messages}
    t0 = log_t0   # keep seconds-from-log-start as the time origin, even after trimming

    bus_health = None
    if args.bus_health:
        baud = meta.get("baudrate") or args.baudrate
        bus_health = analyze_bus_health(frames, meta["dur_sec"], frame_name_by_id,
                                        baudrate=baud, gap_threshold_sec=args.gap_threshold)
        bus_health["baudrate"] = baud
        print(f"Bus health        : avg {bus_health['load_avg']:.2f}% / "
              f"peak {bus_health['load_peak']:.2f}% load (assumed {baud/1000:.0f} kbit/s)")
        if bus_health["bus_off"]:
            print(f"  BUS-OFF/Power-OFF : {len(bus_health['bus_off'])} gap(s) > "
                  f"{args.gap_threshold:g}s with no frames of any ID:")
            for g in bus_health["bus_off"]:
                print(f"    {g['start_clock']} -> {g['end_clock']}   ({g['duration']:.2f}s)")
        else:
            print(f"  BUS-OFF/Power-OFF : none (threshold {args.gap_threshold:g}s)")
        print()

    plot_paths = []
    if not args.no_plots:
        print("Generating graphs…")
        plot_paths = make_plots(series, args.out, t0, show=args.show)
        if plot_paths:
            print(f"  wrote {len([p for p in plot_paths if not os.path.basename(p).startswith('_')])} "
                  f"signal graph(s) + overlays to {os.path.join(args.out, 'plots')}\n")

    if not args.no_excel:
        print("Writing Excel report…")
        xlsx = os.path.join(args.out, os.path.splitext(os.path.basename(args.log))[0] + "_report.xlsx")
        sections = {"summary", "frames", "merged", "charts"}
        if bus_health is not None:
            sections.add("bus_health")
        saved = export_excel(xlsx, meta, [os.path.basename(p) for p in args.dbc],
                             series, frames, frame_name_by_id, args.tol, plot_paths, t0,
                             sections=sections, bus_health=bus_health)
        if saved:
            size_mb = os.path.getsize(saved) / 1e6
            print(f"  saved {saved}  ({size_mb:.1f} MB)\n")

    print(f"Done in {time.time() - t_start:.1f}s.")


def _show_gui_message(kind, title, msg):
    """Best-effort messagebox — there's no console in a windowed build, so
    without this a failure would just look like the app did nothing."""
    try:
        import tkinter as tk
        from tkinter import messagebox
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        getattr(messagebox, kind)(title, msg)
        root.destroy()
    except Exception:
        pass


if __name__ == "__main__":
    # interactive if the user didn't pass both --dbc and --log on the command line
    _flags = ("-h", "--help", "--version")
    _interactive = (("--dbc" not in sys.argv) or ("--log" not in sys.argv)) \
        and not any(f in sys.argv for f in _flags)
    try:
        main()
    except SystemExit as exc:
        if _interactive:
            if exc.code not in (0, None):
                print(exc.code)
                _show_gui_message("showerror", "CAN Log Analyzer", str(exc.code))
        else:
            raise
    except Exception as exc:                # keep the console open to show the traceback
        import traceback
        traceback.print_exc()
        if _interactive:
            _show_gui_message("showerror", "CAN Log Analyzer — error", f"{type(exc).__name__}: {exc}")
    finally:
        # pause so a double-clicked window doesn't vanish — but only when we're
        # actually attached to a console (never when piped / launched headless)
        if _interactive and sys.stdin and sys.stdin.isatty():
            try:
                input("\nPress Enter to close…")
            except EOFError:
                pass
